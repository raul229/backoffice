import re
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from .datos_entel import HC_ARCHIVO_POR_PLAN, HC_COORDS_POR_PLAN
from .mapas import mapa_a_campos

CELDA_RE = re.compile(r"^[A-Z]{1,3}[1-9]\d{0,6}$")
MAX_FILAS_VISTA = 250
MAX_COLS_VISTA = 50
SCAN_FILAS = 400
SCAN_COLS = 60
TEMA_OFFICE = [
    "000000",
    "FFFFFF",
    "EEECE1",
    "1F497D",
    "4F81BD",
    "C0504D",
    "9BBB59",
    "8064A2",
    "4BACC6",
    "F79646",
]


def es_excel(documento) -> bool:
    acepta = documento.acepta or []
    return ".xlsx" in acepta or ".xlsm" in acepta


def normalizar_celda(valor: str) -> str:
    texto = str(valor or "").strip().upper().replace("$", "")
    if not CELDA_RE.match(texto):
        return ""
    return texto


def nombres_hojas(contenido: bytes, keep_vba: bool = False) -> list[str]:
    wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def campos_excel(documento) -> list[dict]:
    campos = [
        {
            **item,
            "celda": normalizar_celda(item.get("celda")),
            "hoja": str(item.get("hoja") or "").strip(),
        }
        for item in mapa_a_campos(documento.mapa or {})
    ]
    campos = [item for item in campos if item["celda"]]
    if campos:
        return campos
    for plan, archivo in HC_ARCHIVO_POR_PLAN.items():
        if archivo != (documento.archivo or ""):
            continue
        return [
            {
                "nombre": campo,
                "hoja": hoja,
                "celda": celda,
                "field": campo,
                "value": "",
                "transform": "",
            }
            for campo, (hoja, celda) in HC_COORDS_POR_PLAN.get(plan, {}).items()
        ]
    return []


def vista_excel(contenido: bytes, campos=None) -> dict:
    wb = load_workbook(BytesIO(contenido), data_only=False, keep_vba=False)
    try:
        por_hoja = {}
        for item in campos or []:
            hoja = str(item.get("hoja") or "").strip()
            celda = normalizar_celda(item.get("celda"))
            if not hoja or not celda:
                continue
            por_hoja.setdefault(hoja, []).append(celda)
        hojas = []
        for nombre in wb.sheetnames:
            hojas.append(_vista_hoja(wb[nombre], nombre, por_hoja.get(nombre, [])))
        return {"hojas": hojas}
    finally:
        wb.close()


def _aplicar_tint(hex6: str, tint: float) -> str:
    if not tint:
        return hex6
    canales = [int(hex6[i : i + 2], 16) for i in (0, 2, 4)]
    out = []
    for canal in canales:
        if tint < 0:
            valor = int(canal * (1 + tint))
        else:
            valor = int(canal + (255 - canal) * tint)
        out.append(max(0, min(255, valor)))
    return "".join(f"{valor:02X}" for valor in out)


def _color_css(color) -> str | None:
    if color is None:
        return None
    tipo = getattr(color, "type", None)
    rgb = getattr(color, "rgb", None)
    if rgb not in (None, "00000000", "0", 0):
        texto = str(rgb)
        hex6 = texto[-6:] if len(texto) >= 6 else ""
        if len(hex6) == 6 and all(c in "0123456789ABCDEFabcdef" for c in hex6):
            if tipo == "rgb" or hex6.upper() not in {"000000"}:
                return f"#{hex6.upper()}"
    if tipo == "theme":
        indice = int(getattr(color, "theme", 0) or 0)
        if 0 <= indice < len(TEMA_OFFICE):
            tint = float(getattr(color, "tint", 0) or 0)
            return f"#{_aplicar_tint(TEMA_OFFICE[indice], tint)}"
    if tipo == "indexed":
        indice = int(getattr(color, "indexed", 0) or 0)
        if 0 <= indice < len(COLOR_INDEX):
            valor = COLOR_INDEX[indice]
            texto = str(valor)
            if len(texto) >= 6:
                return f"#{texto[-6:].upper()}"
    return None


def _estilo_celda(cell) -> dict | None:
    estilo = {}
    fill = cell.fill
    if fill is not None and fill.patternType and fill.patternType != "none":
        fondo = (
            _color_css(getattr(fill, "fgColor", None))
            or _color_css(getattr(fill, "start_color", None))
            or _color_css(getattr(fill, "bgColor", None))
        )
        if fondo:
            estilo["fondo"] = fondo
    font = cell.font
    if font is not None:
        if font.bold:
            estilo["negrita"] = True
        if font.italic:
            estilo["italic"] = True
        color = _color_css(font.color)
        if color:
            estilo["color"] = color
    align = cell.alignment
    if align is not None:
        if align.horizontal in ("center", "centerContinuous"):
            estilo["alinear"] = "center"
        elif align.horizontal == "right":
            estilo["alinear"] = "right"
        if align.wrap_text:
            estilo["wrap"] = True
    border = cell.border
    if border is not None and any(
        getattr(getattr(border, lado, None), "style", None)
        for lado in ("left", "right", "top", "bottom")
    ):
        estilo["borde"] = True
    return estilo or None


def _combinadas(ws, max_fila: int, max_col: int) -> list[dict]:
    items = []
    for rng in ws.merged_cells.ranges:
        if rng.min_row > max_fila or rng.min_col > max_col:
            continue
        max_r = min(rng.max_row, max_fila)
        max_c = min(rng.max_col, max_col)
        filas = max_r - rng.min_row + 1
        cols = max_c - rng.min_col + 1
        if filas < 1 or cols < 1:
            continue
        items.append(
            {
                "celda": f"{get_column_letter(rng.min_col)}{rng.min_row}",
                "filas": filas,
                "cols": cols,
            }
        )
    return items


def _vista_hoja(ws, nombre: str, celdas_mapa: list[str]) -> dict:
    filas_mapa = []
    cols_mapa = []
    for celda in celdas_mapa:
        col, fila = coordinate_from_string(celda)
        filas_mapa.append(fila)
        cols_mapa.append(column_index_from_string(col))
    scan_filas = min(max(ws.max_row or 1, max(filas_mapa, default=1)), SCAN_FILAS)
    scan_cols = min(max(ws.max_column or 1, max(cols_mapa, default=1)), SCAN_COLS)
    valores = {}
    estilos = {}
    max_fila = 1
    max_col = 1
    for row in ws.iter_rows(min_row=1, max_row=scan_filas, min_col=1, max_col=scan_cols):
        for cell in row:
            estilo = _estilo_celda(cell)
            if cell.value not in (None, ""):
                valores[cell.coordinate] = str(cell.value)
                max_fila = max(max_fila, cell.row)
                max_col = max(max_col, cell.column)
            if estilo:
                estilos[cell.coordinate] = estilo
                max_fila = max(max_fila, cell.row)
                max_col = max(max_col, cell.column)
    if filas_mapa:
        max_fila = max(max_fila, max(filas_mapa))
    if cols_mapa:
        max_col = max(max_col, max(cols_mapa))
    for rng in ws.merged_cells.ranges:
        max_fila = max(max_fila, min(rng.max_row, SCAN_FILAS))
        max_col = max(max_col, min(rng.max_col, SCAN_COLS))
    max_fila_vista = min(max_fila + 2, MAX_FILAS_VISTA)
    max_col_vista = min(max_col + 2, MAX_COLS_VISTA)
    if filas_mapa:
        max_fila_vista = max(max_fila_vista, max(filas_mapa))
    if cols_mapa:
        max_col_vista = max(max_col_vista, max(cols_mapa))

    def _en_vista(coord: str) -> bool:
        col, fila = coordinate_from_string(coord)
        return fila <= max_fila_vista and column_index_from_string(col) <= max_col_vista

    anchos = {}
    for indice in range(1, max_col_vista + 1):
        letra = get_column_letter(indice)
        dim = ws.column_dimensions.get(letra)
        ancho = dim.width if dim is not None and dim.width else 10
        anchos[letra] = max(36, int(float(ancho) * 8))
    altos = {}
    for fila in range(1, max_fila_vista + 1):
        dim = ws.row_dimensions.get(fila)
        if dim is not None and dim.height:
            altos[str(fila)] = max(18, int(float(dim.height) * 96 / 72))
    return {
        "nombre": nombre,
        "min_fila": 1,
        "max_fila": max_fila_vista,
        "min_col": 1,
        "max_col": max_col_vista,
        "columnas": [get_column_letter(n) for n in range(1, max_col_vista + 1)],
        "filas": list(range(1, max_fila_vista + 1)),
        "valores": {coord: valor for coord, valor in valores.items() if _en_vista(coord)},
        "estilos": {coord: estilo for coord, estilo in estilos.items() if _en_vista(coord)},
        "combinadas": _combinadas(ws, max_fila_vista, max_col_vista),
        "anchos": anchos,
        "altos": altos,
    }
