"""Llenado de plantillas OIT y Tarifas para Pack Empresas con teléfono fijo."""

from __future__ import annotations

import re
import shutil
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

ASESOR_PACK = "LUIS ZAMBRANO"
GERENTE_PACK = "ANGELA CHONYEN"
PLAZO_CONTRATO = "6 meses"
MONEDA_SOLES = "Soles (S/.)"
PROCESO_ESTADO = "No"
TIPO_DIRECCION = "Instalación"
TIPO_PLAN_TARIFAS = "Otros"
PLAN_TELEFONIA_FIJA = "Pack Empresas_SIM"
BOLSA_TELEFONIA_FIJA = "Ilimitado"
PLAN_TELEFONIA_MOVIL = "Telefonía Móvil_Pack Empresas_SIM"
BOLSA_TELEFONIA_MOVIL = 300
TARIFA_FIJA_CON_IGV = 5.3
TARIFA_MOVIL_CON_IGV = 4.6
TOTAL_BOLSA_TELEFONIA_CON_IGV = TARIFA_FIJA_CON_IGV + TARIFA_MOVIL_CON_IGV  # 9.9 → G33
IGV = Decimal("1.18")

BLOQUEOS_OIT = {
    "H22": "No",  # móvil
    "I22": "Si",  # LDN
    "J22": "Si",  # LDI
    "K22": "No",  # total
    "L22": "Si",  # centrex
}

SUPLEMENTARIOS = (
    ("Bloqueo Larga Distancia Nacional 1era vez", 1, 0),
    ("Bloqueo Larga Distancia Internacional 1era vez", 1, 0),
    ("Instalacion GPON", 1, 0),
)


def renta_sin_igv(precio) -> float:
    valor = Decimal(str(precio or 0)) / IGV
    return float(valor.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _fecha_excel(valor) -> datetime | None:
    if isinstance(valor, datetime):
        return valor
    texto = str(valor or "").strip()
    if not texto:
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def _mes_venta_desde_lista(wb, fecha: datetime | None):
    """Devuelve el datetime exacto de Hoja1 (lista validada de J14) para el mes de la fecha."""
    if fecha is None:
        return None
    hoja = wb["Hoja1"]
    for fila in range(8, 68):
        valor = hoja[f"B{fila}"].value
        if (
            isinstance(valor, datetime)
            and valor.year == fecha.year
            and valor.month == fecha.month
        ):
            return valor
    return datetime(fecha.year, fecha.month, 1)


def _restaurar_validaciones_oit(ws) -> None:
    """openpyxl elimina las data validations x14 al leer; se recrean en formato estándar."""
    existentes = {str(dv.sqref) for dv in ws.data_validations.dataValidation}
    reglas = (
        ("Hoja1!$B$4:$B$5", "J12"),
        ("Hoja1!$B$8:$B$67", "J14"),
        ("Hoja1!$C$8:$C$9", "J15"),
    )
    for formula, celda in reglas:
        if any(celda in str(ref) for ref in existentes):
            continue
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            showInputMessage=True,
        )
        dv.add(celda)
        ws.add_data_validation(dv)


def _solo_digitos(valor: str) -> str:
    return re.sub(r"\D", "", str(valor or ""))


def _numero_telefonico_oit(valor: str) -> int | None:
    """ANI local: quita prefijo 01 de Lima si viene en el N° fijo."""
    digits = _solo_digitos(valor)
    if digits.startswith("01") and len(digits) > 8:
        digits = digits[2:]
    return int(digits) if digits else None


def _consideraciones(contexto: dict) -> str:
    rrll = (contexto.get("RRLL") or "").strip() or "—"
    celular = (contexto.get("CELULAR_RRLL") or "").strip() or "—"
    return (
        f"[Nombre del administrador/back Dealer: {ASESOR_PACK} / "
        f"Gerente:  {GERENTE_PACK}/ "
        f"Contacto de instalación: {rrll} / "
        f"Celular: {celular}"
    )


def llenar_creacion_oit(origen: Path, destino: Path, contexto: dict) -> None:
    shutil.copy(origen, destino)
    wb = load_workbook(destino)
    ws = wb["CREACIÓN OIT"]
    fecha = _fecha_excel(contexto.get("FECHA"))

    ws["E8"] = contexto.get("SAF") or None
    ws["E9"] = contexto.get("RRLL") or None
    ws["J9"] = contexto.get("SIRO") or None
    ws["E11"] = fecha
    ws["J11"] = PLAZO_CONTRATO
    ws["E12"] = fecha
    ws["J12"] = MONEDA_SOLES
    ws["E13"] = None
    precio = contexto.get("PRECIO_PRODUCTO")
    ws["J13"] = renta_sin_igv(precio) if precio not in (None, "") else None
    ws["J14"] = _mes_venta_desde_lista(wb, fecha)
    ws["J15"] = PROCESO_ESTADO

    numero = _numero_telefonico_oit(contexto.get("NUMERO_FIJO") or "")
    ws["D22"] = numero
    for celda, valor in BLOQUEOS_OIT.items():
        ws[celda] = valor
    ws["C30"] = _consideraciones(contexto)

    _restaurar_validaciones_oit(ws)
    wb.save(destino)
    wb.close()


def llenar_tarifas_servicios(origen: Path, destino: Path, contexto: dict) -> None:
    shutil.copy(origen, destino)
    wb = load_workbook(destino, keep_vba=True)
    ws = wb["SOLICITUD T Y S"]

    plan = contexto.get("NOMBRE_PLAN") or "Pack Empresas"
    velocidad = contexto.get("VELOCIDAD") or ""
    precio = contexto.get("PRECIO_PRODUCTO")
    try:
        precio_igv = float(precio) if precio not in (None, "") else None
    except (TypeError, ValueError):
        precio_igv = None

    ws["D10"] = contexto.get("RAZON_SOCIAL") or None
    ws["D11"] = contexto.get("RRLL") or None
    ws["D12"] = ASESOR_PACK
    ws["D16"] = TIPO_DIRECCION
    ws["D17"] = contexto.get("DOMICILIO_INSTALACION") or None
    ws["D19"] = contexto.get("DEPARTAMENTO") or "LIMA"
    ws["D20"] = contexto.get("PROVINCIA") or "LIMA"
    ws["D21"] = contexto.get("DISTRITO") or None
    ws["D26"] = TIPO_PLAN_TARIFAS

    ws["D29"] = "<<opcional>>"
    ws["E29"] = "<<opcional>>"
    ws["G29"] = 0
    ws["D30"] = "<<opcional>>"
    ws["E30"] = "<<opcional>>"
    ws["G30"] = 0

    ws["D31"] = PLAN_TELEFONIA_FIJA
    ws["E31"] = BOLSA_TELEFONIA_FIJA
    ws["G31"] = TARIFA_FIJA_CON_IGV
    ws["D32"] = PLAN_TELEFONIA_MOVIL
    ws["E32"] = BOLSA_TELEFONIA_MOVIL
    ws["G32"] = TARIFA_MOVIL_CON_IGV

    ws["C38"] = f"{plan} {velocidad}".strip()
    # Internet = precio producto CON IGV menos la bolsa de telefonía (G33 = 9.9).
    if precio_igv is not None:
        ws["G38"] = round(precio_igv - TOTAL_BOLSA_TELEFONIA_CON_IGV, 2)
    else:
        ws["G38"] = None

    for indice, (producto, cantidad, tarifa) in enumerate(SUPLEMENTARIOS):
        fila = 44 + indice
        ws[f"D{fila}"] = producto
        ws[f"E{fila}"] = cantidad
        ws[f"F{fila}"] = tarifa

    wb.save(destino)
    wb.close()
