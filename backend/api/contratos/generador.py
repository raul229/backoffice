import io
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

import pymupdf
from django.conf import settings
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from .almacen import bytes_plantilla
from .datos_entel import (
    ARCHIVOS_BASE_POR_PLAN,
    HC_ARCHIVO_POR_PLAN,
    HC_COORDS_POR_PLAN,
    MAPA_ENTEL_POR_ARCHIVO,
    NOMBRE_PNG_POR_PAGINA,
    PLANTILLA_CORREO_ENTEL,
    PLANTILLA_CREACION_OIT,
    PLANTILLA_TARIFAS_SERVICIOS,
)
from .firmas_escan import firmas_a_png, insertar_firmas_en_pdf
from .mapeo import (
    contexto_desde_venta,
    nombre_zip_contrato,
    parsear_direccion_contrato,
    parsear_fecha_contrato,
)
from .plantillas_venta import archivos_pdf_para_venta
from .pack_fijo import llenar_creacion_oit, llenar_tarifas_servicios
from .utils import dividir_texto, generar_eml, partir_en_lineas

_SUFIJO_CAMPO = re.compile(r"_\d+$")


def generar_zip_entel(venta, fecha=None, direccion=None, firmas_paginas=None) -> tuple[bytes, str]:
    contexto, plan, velocidad, promocion = contexto_desde_venta(
        venta,
        fecha=parsear_fecha_contrato(fecha),
        direccion=parsear_direccion_contrato(direccion),
    )
    datos_plan = {
        "_PROMOCION": contexto.get("PROMOCION") or promocion,
        "_RENTA_FIJA": contexto.get("RENTA_FIJA") or "",
        "_DESCUENTO": contexto.get("DESCUENTO") or "",
        "_DESCUENTO_CON_IGV": contexto.get("DESCUENTO_CON_IGV") or "",
        "_NOMBRE_PLAN": contexto.get("NOMBRE_PLAN") or plan,
        "_VELOCIDAD": velocidad,
        "PROMOCION": contexto.get("PROMOCION") or promocion,
        "RENTA_FIJA": contexto.get("RENTA_FIJA") or "",
        "DESCUENTO": contexto.get("DESCUENTO") or "",
        "DESCUENTO_CON_IGV": contexto.get("DESCUENTO_CON_IGV") or "",
        "NOMBRE_PLAN": contexto.get("NOMBRE_PLAN") or plan,
        "VELOCIDAD": velocidad,
    }
    firmas_png = firmas_a_png(firmas_paginas) if firmas_paginas else None
    with tempfile.TemporaryDirectory() as tmp:
        raiz = Path(tmp)
        origenes = raiz / "origen"
        destino = raiz / "salida"
        origenes.mkdir()
        destino.mkdir()
        pdfs, mapas = _materializar_pdfs(origenes, venta, plan, velocidad, promocion)
        for ruta in pdfs:
            firmas = firmas_png if ruta.name in ARCHIVOS_BASE_POR_PLAN.values() else None
            _llenar_pdf(
                ruta,
                destino / ruta.name,
                contexto,
                datos_plan,
                firmas_paginas=firmas,
                mapa=mapas.get(ruta.name),
            )
        if firmas_png:
            for indice, png in firmas_png.items():
                nombre = NOMBRE_PNG_POR_PAGINA.get(indice)
                if nombre:
                    (destino / nombre).write_bytes(png)
        hc_nombre = HC_ARCHIVO_POR_PLAN[plan]
        hc_bytes = bytes_plantilla(hc_nombre, requerida=False)
        if hc_bytes:
            origen_hc = origenes / hc_nombre
            origen_hc.write_bytes(hc_bytes)
            _llenar_hc(origen_hc, destino / hc_nombre, contexto, HC_COORDS_POR_PLAN.get(plan, {}))
        if plan == "Pack Empresas" and str(contexto.get("NUMERO_FIJO") or "").strip():
            _materializar_pack_fijo(origenes, destino, contexto)
        eml_bytes = bytes_plantilla(PLANTILLA_CORREO_ENTEL, requerida=False)
        if eml_bytes:
            extra = {
                "CORREO_BACKOFFICE": getattr(settings, "CORREO_BACKOFFICE", "") or "",
                "CORREO_GERENTE": getattr(settings, "CORREO_GERENTE", "") or "",
                "CORREOS_ADICIONALES": getattr(settings, "CORREOS_ADICIONALES", "") or "",
            }
            origen_eml = origenes / PLANTILLA_CORREO_ENTEL
            origen_eml.write_bytes(eml_bytes)
            adjuntos = [destino / ruta.name for ruta in pdfs]
            generar_eml(origen_eml, {**contexto, **datos_plan, **extra}, adjuntos, destino)
        return _zip_carpeta(destino), nombre_zip_contrato(
            contexto["RAZON_SOCIAL"], contexto["RUC"]
        )


def _materializar_pack_fijo(origenes: Path, destino: Path, contexto: dict) -> None:
    oit_bytes = bytes_plantilla(PLANTILLA_CREACION_OIT, requerida=True)
    origen = origenes / PLANTILLA_CREACION_OIT
    origen.write_bytes(oit_bytes)
    llenar_creacion_oit(origen, destino / PLANTILLA_CREACION_OIT, contexto)
    tarifas_bytes = bytes_plantilla(PLANTILLA_TARIFAS_SERVICIOS, requerida=True)
    origen = origenes / PLANTILLA_TARIFAS_SERVICIOS
    origen.write_bytes(tarifas_bytes)
    llenar_tarifas_servicios(origen, destino / PLANTILLA_TARIFAS_SERVICIOS, contexto)


def _materializar_pdfs(origenes: Path, venta, plan: str, velocidad: int, promocion: str):
    pares = archivos_pdf_para_venta(venta, plan, velocidad, promocion)
    rutas = []
    mapas = {}
    for nombre, mapa in pares:
        ruta = origenes / nombre
        ruta.write_bytes(bytes_plantilla(nombre))
        rutas.append(ruta)
        if mapa:
            mapas[nombre] = mapa
    return rutas, mapas


def _nombre_fuente(campo: str, opciones: dict, contexto: dict, datos_plan: dict):
    nombre = opciones.get("field") or campo
    if nombre in contexto or nombre in datos_plan:
        return nombre
    base = _SUFIJO_CAMPO.sub("", str(nombre))
    if base in contexto or base in datos_plan:
        return base
    return nombre


def _clave_grupo(campo: str, opciones: dict, contexto: dict, datos_plan: dict):
    if "value" in opciones:
        return ("value", campo)
    if "transform" in opciones:
        return ("transform", opciones["transform"])
    return ("ctx", _nombre_fuente(campo, opciones, contexto, datos_plan))


def _valor(campo: str, opciones: dict, contexto: dict, datos_plan: dict):
    if "value" in opciones:
        return opciones["value"]
    if "transform" in opciones:
        return _transform(opciones["transform"], contexto, datos_plan)
    nombre = _nombre_fuente(campo, opciones, contexto, datos_plan)
    if nombre in contexto:
        return contexto[nombre]
    if nombre in datos_plan:
        return datos_plan[nombre]
    return None


def textos_pagina(coords: dict, contexto: dict, datos_plan: dict) -> list[tuple]:
    """(campo, x, y, tamano, texto) ya partido si el mismo dato está varias veces."""
    items = []
    for campo, (x, y, opciones) in (coords or {}).items():
        items.append((campo, x, y, opciones or {}, _valor(campo, opciones or {}, contexto, datos_plan)))

    grupos = {}
    for indice, (campo, _x, _y, opciones, valor) in enumerate(items):
        if valor in (None, "") or not opciones.get("max_caracteres"):
            continue
        grupos.setdefault(_clave_grupo(campo, opciones, contexto, datos_plan), []).append(indice)

    asignados = {}
    resultado = []
    for indice, (campo, x, y, opciones, valor) in enumerate(items):
        if valor in (None, ""):
            continue
        tamano = opciones.get("tamano", 12)
        clave = _clave_grupo(campo, opciones, contexto, datos_plan)
        hermanos = grupos.get(clave, [])
        if indice in asignados:
            texto = asignados[indice]
        elif len(hermanos) > 1:
            limites = [items[i][3].get("max_caracteres") for i in hermanos]
            lineas = partir_en_lineas(str(valor), limites)
            for hermano, linea in zip(hermanos, lineas):
                asignados[hermano] = linea
            texto = asignados[indice]
        elif opciones.get("max_caracteres"):
            texto, _lineas = dividir_texto(str(valor), opciones["max_caracteres"])
        else:
            texto = str(valor)
        if texto in (None, ""):
            continue
        y_inicial = y - (tamano * str(texto).count("\n"))
        resultado.append((campo, x, y_inicial, tamano, texto))
    return resultado


def _transform(nombre: str, contexto: dict, datos_plan: dict) -> str | None:
    if nombre == "doc_y_dni":
        return f"{contexto.get('TIPO_DOCUMENTO_RRLL', '')} {contexto.get('DNI', '')}".strip()
    if nombre == "velocidades":
        v = int(contexto.get("VELOCIDAD") or datos_plan.get("_VELOCIDAD") or 0)
        con_bono = int(contexto.get("VELOCIDAD_CON_BONO") or v)
        return f"{v}                {con_bono}               {v}             {con_bono}"
    if nombre == "nombre_plan":
        return f"{contexto.get('NOMBRE_PLAN') or datos_plan.get('_NOMBRE_PLAN')} {datos_plan.get('_VELOCIDAD')}"
    if nombre == "dia_mes_anio":
        return (
            f"{contexto.get('DIA', '')}                                    "
            f"{contexto.get('NOMBRE_MES', '')}                            "
            f"{contexto.get('ANIO', '')}"
        )
    if nombre == "descuento":
        return f"S/. {contexto.get('DESCUENTO') or datos_plan.get('_DESCUENTO')}"
    if nombre == "descuento_con_igv":
        return f"S/. {contexto.get('DESCUENTO_CON_IGV') or datos_plan.get('_DESCUENTO_CON_IGV')}"
    if nombre == "precio_int_solo":
        bruto = str(contexto.get("RENTA_FIJA") or datos_plan.get("_RENTA_FIJA") or "0")
        renta = float(bruto.replace("S/.", "").strip() or "0")
        return f"S/. {renta - 9.9:.2f}"
    return None


def _llenar_pdf(
    origen: Path,
    destino: Path,
    contexto: dict,
    datos_plan: dict,
    firmas_paginas: dict[int, bytes] | None = None,
    mapa: dict | None = None,
) -> None:
    pdf = pymupdf.open(origen)
    try:
        paginas = mapa if mapa else MAPA_ENTEL_POR_ARCHIVO.get(origen.name, {})
        for num_pagina, coords in paginas.items():
            if num_pagina >= len(pdf):
                raise ValidationError(
                    {
                        "detail": (
                            f"El PDF {origen.name} no tiene la página {num_pagina} "
                            f"(tiene {len(pdf)})."
                        )
                    }
                )
            pagina = pdf[num_pagina]
            for _campo, x, y, tamano, texto in textos_pagina(coords, contexto, datos_plan):
                pagina.insert_text((x, y), texto, fontsize=tamano)
        # Firmas manuscritas encima del texto, en la zona de firma (no reemplazan la hoja).
        if firmas_paginas:
            insertar_firmas_en_pdf(pdf, firmas_paginas)
        pdf.save(str(destino))
    finally:
        pdf.close()


def _llenar_hc(origen: Path, destino: Path, contexto: dict, coords: dict) -> None:
    shutil.copy(origen, destino)
    wb = load_workbook(destino, keep_vba=True)
    for campo, (hoja, celda) in coords.items():
        valor = contexto.get(campo)
        if valor in (None, ""):
            continue
        wb[hoja][celda] = valor
    for nombre in wb.sheetnames:
        wb[nombre].sheet_state = "visible" if nombre == "Formulario" else "hidden"
    wb.save(destino)


def _zip_carpeta(carpeta: Path) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archivo:
        for ruta in sorted(carpeta.iterdir()):
            if ruta.is_file():
                archivo.write(ruta, ruta.name)
    return buffer.getvalue()
