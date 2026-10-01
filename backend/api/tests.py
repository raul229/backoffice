from django.contrib.auth.models import Group, User
from rest_framework import status
from rest_framework.test import APITestCase
from unittest.mock import patch
from pathlib import Path

from .models import (
    Cliente,
    Direccion,
    Empresa,
    Flujo,
    FlujoPaso,
    Paso,
    Persona,
    PlantillaContrato,
    Producto,
    Promocion,
    TipoCliente,
    Venta,
)

# Solo para tests: origen opcional de bytes al sembrar la BD.
_FIXTURES_PLANTILLAS = Path(__file__).resolve().parent / "fixtures_plantillas"
_PLANTILLAS_DISCO_TEST = Path("/home/raul/Proyectos/Python/Contratos/CONTRATOS_ENTEL")


class ApiEndpointsTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("tester", "tester@test.com", "pass")
        self.client.force_authenticate(self.user)

    def _cargar_plantillas_en_bd(self, requeridas=None):
        """Siembra PlantillaContrato desde fixtures de prueba (no es runtime)."""
        from .contratos.datos_entel import CATALOGO_PLANTILLAS_ENTEL

        cargadas = 0
        for item in CATALOGO_PLANTILLAS_ENTEL:
            if requeridas is not None and item["clave"] not in requeridas:
                continue
            ruta = None
            for base in (_FIXTURES_PLANTILLAS, _PLANTILLAS_DISCO_TEST):
                candidata = base / item["archivo"]
                if candidata.is_file():
                    ruta = candidata
                    break
            if ruta is None:
                continue
            data = ruta.read_bytes()
            PlantillaContrato.objects.update_or_create(
                clave=item["clave"],
                defaults={
                    "nombre_archivo": item["archivo"],
                    "content_type": "application/octet-stream",
                    "tamano": len(data),
                    "contenido": data,
                    "actualizado_por": self.user,
                },
            )
            cargadas += 1
        if cargadas == 0:
            self.skipTest("No hay archivos de plantilla para sembrar en BD de prueba")
        return cargadas

    def test_choices_endpoint_returns_frontend_options(self):
        response = self.client.get("/api/choices/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("tipos_cliente", response.data)
        self.assertIn({"value": TipoCliente.PERSONA, "label": "Persona Natural"}, response.data["tipos_cliente"])

    def test_create_persona_creates_cliente_when_missing(self):
        response = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "DNI",
                "numero_documento": "12345678",
                "nombres": "Ana",
                "apellidos": "Perez",
                "distrito_nacimiento": "Lima",
                "padre": "Carlos",
                "madre": "Maria",
                "celular": "987654321",
                "correo": "Ana.Perez@Example.COM",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        cliente = Cliente.objects.get(id=response.data["cliente"])
        self.assertEqual(cliente.tipo, TipoCliente.PERSONA)
        self.assertEqual(cliente.correo, "ana.perez@example.com")

        missing = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "DNI",
                "numero_documento": "12345679",
                "nombres": "Ana",
                "apellidos": "Perez",
                "celular": "987654321",
            },
            format="json",
        )
        self.assertEqual(missing.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("correo", missing.data)

    def test_update_persona_parents_and_birth(self):
        created = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "DNI",
                "numero_documento": "12345670",
                "nombres": "Ana",
                "apellidos": "Perez",
                "celular": "987654321",
                "correo": "ana.perez@example.com",
            },
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertEqual(created.data["padre"], "")
        self.assertEqual(created.data["madre"], "")
        self.assertEqual(created.data["distrito_nacimiento"], "")

        response = self.client.patch(
            f"/api/personas/{created.data['id']}/",
            {
                "padre": "carlos perez",
                "madre": "maria lopez",
                "distrito_nacimiento": "lince",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["padre"], "CARLOS PEREZ")
        self.assertEqual(response.data["madre"], "MARIA LOPEZ")
        self.assertEqual(response.data["distrito_nacimiento"], "LINCE")
        self.assertEqual(response.data["nombres"], "ANA")
        cliente = self.client.get(f"/api/clientes/{created.data['cliente']}/")
        self.assertEqual(cliente.data["correo"], "ana.perez@example.com")
        self.assertEqual(cliente.data["persona"]["padre"], "CARLOS PEREZ")

    def test_create_persona_carnet_extranjeria_requires_nine_digits(self):
        response = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "CE",
                "numero_documento": "12345678",
                "nombres": "Luis",
                "apellidos": "Diaz",
                "celular": "987654321",
                "correo": "ana.perez@example.com",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        response = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "CE",
                "numero_documento": "123456789",
                "nombres": "Luis",
                "apellidos": "Diaz",
                "celular": "987654321",
                "correo": "ana.perez@example.com",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

    def test_create_venta_generates_workflow_steps(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 500", velocidad=500, precio=99)
        promocion = Promocion.objects.create(nombre="Promo", descripcion="Instalacion gratis")
        flujo = Flujo.objects.create(nombre="Instalacion", tipo_cliente=TipoCliente.PERSONA)
        paso_uno = Paso.objects.create(nombre="Validacion", descripcion="Validar datos")
        paso_dos = Paso.objects.create(nombre="Instalacion", descripcion="Instalar servicio")
        FlujoPaso.objects.create(flujo=flujo, paso=paso_uno, orden=1)
        FlujoPaso.objects.create(flujo=flujo, paso=paso_dos, orden=2)

        response = self.client.post(
            "/api/ventas/",
            {
                "cliente": cliente.id,
                "producto": producto.id,
                "flujo": flujo.id,
                "promociones": [promocion.id],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(len(response.data["pasos"]), 2)
        self.assertEqual(response.data["promociones"], [promocion.id])
        self.assertIn("actualizado", response.data)

    def test_ventas_list_orders_recently_updated_first(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 60", velocidad=60, precio=45)
        flujo = Flujo.objects.create(nombre="Flujo orden", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        primera = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        segunda = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(primera.status_code, status.HTTP_201_CREATED)
        self.assertEqual(segunda.status_code, status.HTTP_201_CREATED)

        lista = self.client.get("/api/ventas/")
        self.assertEqual(
            [venta["id"] for venta in lista.data],
            [segunda.data["id"], primera.data["id"]],
        )

        patched = self.client.patch(
            f"/api/ventas/{primera.data['id']}/",
            {"psi": "ABC123"},
            format="json",
        )
        self.assertEqual(patched.status_code, status.HTTP_200_OK)
        reordenada = self.client.get("/api/ventas/")
        self.assertEqual(reordenada.data[0]["id"], primera.data["id"])
        self.assertGreater(reordenada.data[0]["actualizado"], primera.data["actualizado"])

    def test_venta_keeps_its_own_direccion(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 100", velocidad=100, precio=70)
        flujo = Flujo.objects.create(nombre="Instalacion", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Validacion", descripcion="Validar")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)
        primera = Direccion.objects.create(
            cliente=cliente,
            tipo="JIRON",
            direccion="PUNO",
            numero="654",
            distrito="HUANCAYO",
        )
        segunda = Direccion.objects.create(
            cliente=cliente,
            tipo="CALLE",
            direccion="LUIGGI BARSATO",
            numero="167",
            distrito="SAN BORJA",
        )

        venta_uno = self.client.post(
            "/api/ventas/",
            {
                "cliente": cliente.id,
                "direccion": primera.id,
                "producto": producto.id,
                "flujo": flujo.id,
            },
            format="json",
        )
        venta_dos = self.client.post(
            "/api/ventas/",
            {
                "cliente": cliente.id,
                "direccion": segunda.id,
                "producto": producto.id,
                "flujo": flujo.id,
            },
            format="json",
        )
        self.assertEqual(venta_uno.status_code, status.HTTP_201_CREATED)
        self.assertEqual(venta_dos.status_code, status.HTTP_201_CREATED)
        self.assertEqual(venta_uno.data["direccion_detalle"]["distrito"], "HUANCAYO")
        self.assertEqual(venta_dos.data["direccion_detalle"]["distrito"], "SAN BORJA")

        otro = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        ajena = Direccion.objects.create(
            cliente=otro,
            tipo="CALLE",
            direccion="OTRA",
            numero="1",
            distrito="LIMA",
        )
        rejected = self.client.post(
            "/api/ventas/",
            {
                "cliente": cliente.id,
                "direccion": ajena.id,
                "producto": producto.id,
                "flujo": flujo.id,
            },
            format="json",
        )
        self.assertEqual(rejected.status_code, status.HTTP_400_BAD_REQUEST)

    def test_create_empresa_creates_cliente_juridico(self):
        persona = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "DNI",
                "numero_documento": "87654321",
                "nombres": "Luis",
                "apellidos": "Rojas",
                "distrito_nacimiento": "Lima",
                "padre": "Pedro",
                "madre": "Ana",
                "celular": "999888777",
                "correo": "luis.rojas@example.com",
            },
            format="json",
        )
        response = self.client.post(
            "/api/empresas/",
            {
                "ruc": "20123456789",
                "razon_social": "Empresa SAC",
                "representante_legal": persona.data["id"],
                "correo": "facturacion@empresa.com",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        cliente = Cliente.objects.get(id=response.data["cliente"])
        self.assertEqual(cliente.tipo, TipoCliente.EMPRESA)
        self.assertEqual(cliente.correo, "facturacion@empresa.com")

    def test_venta_rejects_flujo_of_other_client_type(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.EMPRESA)
        producto = Producto.objects.create(nombre="Fibra 200", velocidad=200, precio=80)
        flujo = Flujo.objects.create(nombre="Flujo ruc 10", tipo_cliente=TipoCliente.PERSONA)

        response = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_update_venta_flujo_rebuilds_steps(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 100", velocidad=100, precio=70)
        flujo_a = Flujo.objects.create(nombre="Flujo A", tipo_cliente=TipoCliente.PERSONA)
        flujo_b = Flujo.objects.create(nombre="Flujo B", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Validacion", descripcion="Validar")
        extra = Paso.objects.create(nombre="Instalacion", descripcion="Instalar")
        FlujoPaso.objects.create(flujo=flujo_a, paso=paso, orden=1)
        FlujoPaso.objects.create(flujo=flujo_b, paso=paso, orden=1)
        FlujoPaso.objects.create(flujo=flujo_b, paso=extra, orden=2)

        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo_a.id},
            format="json",
        )
        self.assertEqual(len(created.data["pasos"]), 1)

        updated = self.client.patch(
            f"/api/ventas/{created.data['id']}/",
            {"flujo": flujo_b.id},
            format="json",
        )
        self.assertEqual(updated.status_code, status.HTTP_200_OK)
        self.assertEqual(len(updated.data["pasos"]), 2)
        self.assertEqual(updated.data["flujo"], flujo_b.id)

    def test_aprobar_todos_los_pasos_marca_venta_instalada(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 100", velocidad=100, precio=70)
        flujo = Flujo.objects.create(nombre="Instalacion", tipo_cliente=TipoCliente.PERSONA)
        paso_uno = Paso.objects.create(nombre="Validacion", descripcion="Validar")
        paso_dos = Paso.objects.create(nombre="Instalacion", descripcion="Instalar")
        FlujoPaso.objects.create(flujo=flujo, paso=paso_uno, orden=1)
        FlujoPaso.objects.create(flujo=flujo, paso=paso_dos, orden=2)

        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertEqual(created.data["estado"], "EN_PROCESO")
        pasos = created.data["pasos"]

        self.client.patch(f"/api/venta-pasos/{pasos[0]['id']}/", {"estado": "APROBADO"}, format="json")
        self.client.patch(f"/api/venta-pasos/{pasos[1]['id']}/", {"estado": "EN_PROCESO"}, format="json")
        venta_abierta = self.client.get(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(venta_abierta.data["estado"], "EN_PROCESO")

        self.client.patch(f"/api/venta-pasos/{pasos[1]['id']}/", {"estado": "APROBADO"}, format="json")
        detalle = self.client.get(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(detalle.status_code, status.HTTP_200_OK)
        self.assertEqual(detalle.data["estado"], "INSTALADO")
        self.assertTrue(all(paso["estado"] == "APROBADO" for paso in detalle.data["pasos"]))

        self.client.patch(f"/api/venta-pasos/{pasos[1]['id']}/", {"estado": "PENDIENTE"}, format="json")
        reabierta = self.client.get(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(reabierta.data["estado"], "EN_PROCESO")

    def test_delete_flujo_paso_used_by_venta(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 50", velocidad=50, precio=40)
        flujo = Flujo.objects.create(nombre="Flujo", tipo_cliente=TipoCliente.PERSONA)
        paso_uno = Paso.objects.create(nombre="Uno", descripcion="Uno")
        paso_dos = Paso.objects.create(nombre="Dos", descripcion="Dos")
        primero = FlujoPaso.objects.create(flujo=flujo, paso=paso_uno, orden=1)
        FlujoPaso.objects.create(flujo=flujo, paso=paso_dos, orden=2)

        self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )

        response = self.client.delete(f"/api/flujo-pasos/{primero.id}/")
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(list(flujo.flujopaso_set.order_by("orden").values_list("orden", "paso__nombre")), [(1, "Dos")])

    def test_add_flujo_paso_syncs_existing_venta(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 80", velocidad=80, precio=55)
        flujo = Flujo.objects.create(nombre="Flujo sync", tipo_cliente=TipoCliente.PERSONA)
        paso_uno = Paso.objects.create(nombre="Uno", descripcion="Uno")
        paso_dos = Paso.objects.create(nombre="Dos", descripcion="Dos")
        FlujoPaso.objects.create(flujo=flujo, paso=paso_uno, orden=1)

        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(len(created.data["pasos"]), 1)

        self.client.post(
            "/api/flujo-pasos/",
            {"flujo": flujo.id, "paso": paso_dos.id, "orden": 2},
            format="json",
        )
        detalle = self.client.get(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(len(detalle.data["pasos"]), 2)
        self.assertEqual(
            [paso["flujo_paso_detalle"]["paso_detalle"]["nombre"] for paso in detalle.data["pasos"]],
            ["Uno", "Dos"],
        )

    def test_reorder_flujo_pasos_keeps_venta_paso_estado(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 90", velocidad=90, precio=60)
        flujo = Flujo.objects.create(nombre="Flujo reorder", tipo_cliente=TipoCliente.PERSONA)
        paso_uno = Paso.objects.create(nombre="Uno", descripcion="Uno")
        paso_dos = Paso.objects.create(nombre="Dos", descripcion="Dos")
        paso_tres = Paso.objects.create(nombre="Tres", descripcion="Tres")
        primero = FlujoPaso.objects.create(flujo=flujo, paso=paso_uno, orden=1)
        segundo = FlujoPaso.objects.create(flujo=flujo, paso=paso_dos, orden=2)

        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        venta_paso_uno = next(
            paso for paso in created.data["pasos"] if paso["flujo_paso"] == primero.id
        )
        self.client.patch(
            f"/api/venta-pasos/{venta_paso_uno['id']}/",
            {"estado": "APROBADO"},
            format="json",
        )

        tercero = self.client.post(
            "/api/flujo-pasos/",
            {"flujo": flujo.id, "paso": paso_tres.id, "orden": 3},
            format="json",
        )
        self.assertEqual(tercero.status_code, status.HTTP_201_CREATED)

        invalid = self.client.patch(
            f"/api/flujos/{flujo.id}/reordenar-pasos/",
            {"ids": [tercero.data["id"], primero.id]},
            format="json",
        )
        self.assertEqual(invalid.status_code, status.HTTP_400_BAD_REQUEST)

        response = self.client.patch(
            f"/api/flujos/{flujo.id}/reordenar-pasos/",
            {"ids": [tercero.data["id"], primero.id, segundo.id]},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            [paso["paso_detalle"]["nombre"] for paso in response.data["pasos_detalle"]],
            ["Tres", "Uno", "Dos"],
        )
        self.assertEqual(
            [paso["orden"] for paso in response.data["pasos_detalle"]],
            [1, 2, 3],
        )

        detalle = self.client.get(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(detalle.status_code, status.HTTP_200_OK)
        nombres = [paso["flujo_paso_detalle"]["paso_detalle"]["nombre"] for paso in detalle.data["pasos"]]
        estados = {
            paso["flujo_paso_detalle"]["paso_detalle"]["nombre"]: paso["estado"]
            for paso in detalle.data["pasos"]
        }
        self.assertEqual(nombres, ["Tres", "Uno", "Dos"])
        self.assertEqual(estados["Uno"], "APROBADO")
        self.assertEqual(estados["Dos"], "PENDIENTE")
        self.assertEqual(estados["Tres"], "PENDIENTE")
        self.assertEqual(
            next(paso["id"] for paso in detalle.data["pasos"] if paso["flujo_paso"] == primero.id),
            venta_paso_uno["id"],
        )
        self.assertEqual(detalle.data["estado"], "EN_PROCESO")

    def test_user_only_sees_own_ventas(self):
        grupo = Group.objects.get(name="Asesor")
        ana = User.objects.create_user("ana_venta", password="secret")
        luis = User.objects.create_user("luis_venta", password="secret")
        ana.groups.add(grupo)
        luis.groups.add(grupo)

        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 10", velocidad=10, precio=20)
        flujo = Flujo.objects.create(nombre="Flujo propio", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        self.client.force_authenticate(ana)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertEqual(created.data["creado_por"], ana.id)

        self.client.force_authenticate(luis)
        lista = self.client.get("/api/ventas/")
        self.assertEqual(lista.status_code, status.HTTP_200_OK)
        self.assertEqual(lista.data, [])
        detalle = self.client.get(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(detalle.status_code, status.HTTP_404_NOT_FOUND)

        supervisor = User.objects.create_user("super_venta", password="secret")
        supervisor.groups.add(Group.objects.get(name="Supervisor"))
        self.client.force_authenticate(supervisor)
        todas = self.client.get("/api/ventas/")
        self.assertEqual(len(todas.data), 1)

    def test_create_and_update_producto_and_promocion(self):
        created = self.client.post(
            "/api/productos/",
            {"nombre": "fibra 300", "velocidad": 300, "precio": "89.90"},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertEqual(created.data["nombre"], "FIBRA 300")
        self.assertEqual(created.data["tipo_cliente"], TipoCliente.PERSONA)

        empresa = self.client.post(
            "/api/productos/",
            {
                "nombre": "fibra empresa",
                "velocidad": 600,
                "precio": "129.90",
                "tipo_cliente": TipoCliente.EMPRESA,
            },
            format="json",
        )
        self.assertEqual(empresa.status_code, status.HTTP_201_CREATED)
        self.assertEqual(empresa.data["tipo_cliente"], TipoCliente.EMPRESA)

        updated = self.client.patch(
            f"/api/productos/{created.data['id']}/",
            {"precio": "95.00"},
            format="json",
        )
        self.assertEqual(updated.status_code, status.HTTP_200_OK)
        self.assertEqual(updated.data["precio"], "95.00")

        promo = self.client.post(
            "/api/promociones/",
            {"nombre": "instalacion gratis", "descripcion": "sin costo el primer mes"},
            format="json",
        )
        self.assertEqual(promo.status_code, status.HTTP_201_CREATED)
        self.assertEqual(promo.data["nombre"], "INSTALACION GRATIS")
        self.assertEqual(promo.data["descripcion"], "SIN COSTO EL PRIMER MES")

        renamed = self.client.patch(
            f"/api/promociones/{promo.data['id']}/",
            {"nombre": "promo 2 meses"},
            format="json",
        )
        self.assertEqual(renamed.status_code, status.HTTP_200_OK)
        self.assertEqual(renamed.data["nombre"], "PROMO 2 MESES")

        deleted = self.client.delete(f"/api/promociones/{promo.data['id']}/")
        self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)

    def test_venta_rejects_producto_of_other_client_type(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.EMPRESA)
        producto = Producto.objects.create(
            nombre="Fibra casa", velocidad=100, precio=70, tipo_cliente=TipoCliente.PERSONA
        )
        flujo = Flujo.objects.create(nombre="Empresa", tipo_cliente=TipoCliente.EMPRESA)
        response = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("producto", response.data)

    def _venta_empresa_entel(self, nombre_producto="Internet Empresas", velocidad=200, promo="bono de velocidad por 6m"):
        rrll = Persona.objects.create(
            cliente=Cliente.objects.create(tipo=TipoCliente.PERSONA),
            tipo_documento="DNI",
            numero_documento="87654321",
            nombres="Ana",
            apellidos="Perez",
            celular="987654321",
        )
        empresa_cliente = Cliente.objects.create(tipo=TipoCliente.EMPRESA, correo="ana@empresa.com")
        Empresa.objects.create(
            cliente=empresa_cliente,
            ruc="20522317285",
            razon_social="INDOTECH SAC",
            representante_legal=rrll,
        )
        direccion = Direccion.objects.create(
            cliente=empresa_cliente,
            tipo="CALLE",
            direccion="Luiggi Barsato",
            numero="167",
            distrito="San Borja",
        )
        producto = Producto.objects.create(
            nombre=nombre_producto,
            velocidad=velocidad,
            precio=90,
            tipo_cliente=TipoCliente.EMPRESA,
        )
        flujo = Flujo.objects.create(nombre="Empresa", tipo_cliente=TipoCliente.EMPRESA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)
        promocion = Promocion.objects.create(nombre=promo, descripcion="Entel")
        created = self.client.post(
            "/api/ventas/",
            {
                "cliente": empresa_cliente.id,
                "producto": producto.id,
                "flujo": flujo.id,
                "direccion": direccion.id,
                "promociones": [promocion.id],
                "psi": "PSI1",
                "siro": "SIRO1",
                "numero_oportunidad": "OPP1",
            },
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        return created.data["id"]

    def test_generar_contrato_rejects_persona_natural(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 200", velocidad=200, precio=90)
        flujo = Flujo.objects.create(nombre="Persona", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        response = self.client.post(f"/api/ventas/{created.data['id']}/generar-contrato/")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("persona jurídica", response.data["detail"])

    def test_asesor_cannot_generar_contrato(self):
        ana = User.objects.create_user("ana_contrato", password="secret")
        ana.groups.add(Group.objects.get(name="Asesor"))
        venta_id = self._venta_empresa_entel()
        self.client.force_authenticate(ana)
        response = self.client.post(f"/api/ventas/{venta_id}/generar-contrato/")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def _jpeg_minimo(self):
        """JPEG pequeño válido para tests de firmas escaneadas."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page(width=200, height=280)
        page.draw_rect(page.rect, color=(0.8, 0.8, 0.8), fill=(0.9, 0.9, 0.9))
        data = page.get_pixmap().tobytes("jpeg")
        doc.close()
        return SimpleUploadedFile("firma.jpg", data, content_type="image/jpeg")

    def _post_generar_contrato(self, venta_id, data=None, con_firmas=True):
        payload = dict(data or {})
        if con_firmas:
            payload["firma_pagina_3"] = self._jpeg_minimo()
            payload["firma_pagina_5"] = self._jpeg_minimo()
            return self.client.post(
                f"/api/ventas/{venta_id}/generar-contrato/",
                payload,
                format="multipart",
            )
        return self.client.post(
            f"/api/ventas/{venta_id}/generar-contrato/",
            payload,
            format="json",
        )

    def test_generar_contrato_entel_returns_zip(self):
        import zipfile
        from io import BytesIO

        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel()
        response = self._post_generar_contrato(venta_id)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/zip")
        self.assertIn("INDOTECH SAC - 20522317285.zip", response["Content-Disposition"])
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            nombres = set(archivo.namelist())
        self.assertIn("internet empresas.pdf", nombres)
        self.assertIn("bono duplica.pdf", nombres)

    def test_generar_contrato_internet_exige_firmas(self):
        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel()
        response = self._post_generar_contrato(venta_id, con_firmas=False)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("firma", response.data["detail"].lower())

    def test_generar_contrato_fecha_invalida(self):
        venta_id = self._venta_empresa_entel()
        response = self._post_generar_contrato(venta_id, {"fecha": "no-es-fecha"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("fecha", response.data)

    def test_generar_contrato_acepta_fecha_pasada(self):
        import pymupdf
        import zipfile
        from io import BytesIO

        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel()
        response = self._post_generar_contrato(venta_id, {"fecha": "2024-03-15"})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            pdf = pymupdf.open(stream=archivo.read("internet empresas.pdf"), filetype="pdf")
            # Páginas de firma reemplazadas por imagen (índices 2 y 4).
            self.assertTrue(pdf[2].get_images())
            self.assertTrue(pdf[4].get_images())
            # Página 1 sigue llevando datos del cliente.
            texto = pdf[0].get_text()
            self.assertIn("INDOTECH", texto.upper())
            # Sobre la foto de firma se vuelve a pintar RRLL y fecha.
            texto_p3 = pdf[2].get_text().upper()
            texto_p5 = pdf[4].get_text().upper()
            self.assertIn("ANA PEREZ", texto_p3)
            self.assertIn("15/03/2024", pdf[2].get_text())
            self.assertIn("ANA PEREZ", texto_p5)
            self.assertIn("987654321", pdf[4].get_text())
            pdf.close()

    def test_generar_contrato_acepta_direccion_personalizada(self):
        import pymupdf
        import zipfile
        from io import BytesIO

        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel()
        direccion = "AV JAVIER PRADO ESTE 4450, SAN ISIDRO - PORTAL FACTIBILIDAD"
        response = self._post_generar_contrato(
            venta_id,
            {"fecha": "2024-03-15", "direccion": direccion},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            texto = "".join(
                page.get_text()
                for page in pymupdf.open(stream=archivo.read("internet empresas.pdf"), filetype="pdf")
            )
        self.assertIn("PORTAL FACTIBILIDAD", texto)
        self.assertNotIn("Luiggi Barsato", texto)

    def test_pack_fijo_llena_oit_y_tarifas(self):
        import zipfile
        from io import BytesIO
        from openpyxl import load_workbook

        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel(
            nombre_producto="Pack Empresas",
            velocidad=300,
            promo="30% y bono de velocidad por 6m",
        )
        venta = Venta.objects.get(id=venta_id)
        venta.numero_fijo = "01-7131046"
        venta.saf = "309341"
        venta.siro = "77422"
        venta.cotizacion = "COT-99"
        venta.save(update_fields=["numero_fijo", "saf", "siro", "cotizacion"])
        venta.producto.precio = "129.90"
        venta.producto.save(update_fields=["precio"])
        response = self.client.post(
            f"/api/ventas/{venta_id}/generar-contrato/",
            {"fecha": "2026-09-21"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            nombres = set(archivo.namelist())
            self.assertIn("plantilla_creacion_oit.xlsx", nombres)
            self.assertIn("plantilla_tarifas_y_servicios.xlsm", nombres)
            oit = load_workbook(BytesIO(archivo.read("plantilla_creacion_oit.xlsx")))[
                "CREACIÓN OIT"
            ]
            self.assertEqual(str(oit["E8"].value), "309341")
            self.assertEqual(oit["J11"].value, "6 meses")
            self.assertEqual(oit["J13"].value, 110.08)
            self.assertEqual(oit["D22"].value, 7131046)
            self.assertIn("LUIS ZAMBRANO", oit["C30"].value)
            tarifas = load_workbook(
                BytesIO(archivo.read("plantilla_tarifas_y_servicios.xlsm")),
                keep_vba=True,
            )["SOLICITUD T Y S"]
            self.assertEqual(tarifas["D12"].value, "LUIS ZAMBRANO")
            self.assertEqual(tarifas["G31"].value, 5.3)
            self.assertEqual(float(tarifas["G38"].value), 120.0)
            self.assertEqual(tarifas["C38"].value, "Pack Empresas 300")

    def test_pack_sin_fijo_no_incluye_oit_tarifas(self):
        import zipfile
        from io import BytesIO

        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel(
            nombre_producto="Pack Empresas",
            velocidad=300,
            promo="Solo 30% por 6m.",
        )
        response = self.client.post(f"/api/ventas/{venta_id}/generar-contrato/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            nombres = set(archivo.namelist())
        self.assertNotIn("plantilla_creacion_oit.xlsx", nombres)
        self.assertNotIn("plantilla_tarifas_y_servicios.xlsm", nombres)

    def test_generar_contrato_sin_plantilla_en_bd(self):
        PlantillaContrato.objects.all().delete()
        venta_id = self._venta_empresa_entel(
            nombre_producto="Pack Empresas",
            velocidad=300,
            promo="Solo 30% por 6m.",
        )
        response = self.client.post(f"/api/ventas/{venta_id}/generar-contrato/")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Configuración", response.data["detail"])


    def test_correo_cuenta_planner_rejects_persona_natural(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 200", velocidad=200, precio=90)
        flujo = Flujo.objects.create(nombre="Persona", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        response = self.client.post(f"/api/ventas/{created.data['id']}/correo-cuenta-planner/")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("persona jurídica", response.data["detail"])

    def test_correo_cuenta_planner_genera_eml(self):
        from unittest.mock import patch
        from email import message_from_bytes

        venta_id = self._venta_empresa_entel()
        with patch(
            "api.contratos.correo_cuenta.consultar_sunat",
            return_value={
                "ruc": "20522317285",
                "razon_social": "INDOTECH SAC",
                "tipo_cliente": "EMPRESA",
                "ficha": [
                    {"label": "Número de RUC", "value": "20522317285 - INDOTECH SAC"},
                    {"label": "Tipo Contribuyente", "value": "SOCIEDAD ANONIMA CERRADA"},
                    {"label": "Estado del Contribuyente", "value": "ACTIVO"},
                    {"label": "Domicilio Fiscal", "value": "CAL. LUIGGI BARSATO NRO. 167 SAN BORJA"},
                ],
                "representantes": [
                    {
                        "tipo_documento": "DNI",
                        "numero_documento": "87654321",
                        "nombre_completo": "PEREZ ANA",
                        "cargo": "GERENTE",
                        "fecha_desde": "01/01/2020",
                    }
                ],
            },
        ):
            response = self.client.post(f"/api/ventas/{venta_id}/correo-cuenta-planner/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(
            "SOLICITUD CREACION DE CUENTA  20522317285- INDOTECH SAC.eml",
            response["Content-Disposition"],
        )
        msg = message_from_bytes(response.content)
        self.assertIn("SOLICITUD CREACION DE CUENTA // 20522317285- INDOTECH SAC", msg["Subject"])
        self.assertIn("asignaciondecuentas@entel.pe", msg["To"])
        cuerpo = response.content.decode("utf-8", errors="replace")
        self.assertIn("Tipo Contribuyente", cuerpo)
        self.assertIn("87654321", cuerpo)
        self.assertIn("ana@empresa.com", cuerpo)
        self.assertIn("987654321", cuerpo)
        self.assertIn("Content-Transfer-Encoding: 8bit", cuerpo)
        self.assertNotIn("quoted-printable", cuerpo)
        self.assertNotIn("=3D", cuerpo)

    def test_asesor_cannot_ver_plantillas_contrato(self):
        ana = User.objects.create_user("ana_plantilla", password="secret")
        ana.groups.add(Group.objects.get(name="Asesor"))
        self.client.force_authenticate(ana)
        response = self.client.get("/api/plantillas-contrato/")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_plantillas_contrato_catalogo_y_carga(self):
        import pymupdf
        from django.core.files.uploadedfile import SimpleUploadedFile

        lista = self.client.get("/api/plantillas-contrato/")
        self.assertEqual(lista.status_code, status.HTTP_200_OK)
        claves = [item["clave"] for item in lista.data]
        self.assertIn("internet-empresas", claves)
        self.assertFalse(lista.data[0]["cargada"])

        pdf = pymupdf.open()
        for _ in range(5):
            pdf.new_page()
        contenido = pdf.tobytes()
        pdf.close()
        subida = self.client.post(
            "/api/plantillas-contrato/internet-empresas/",
            {"archivo": SimpleUploadedFile("contrato.pdf", contenido, content_type="application/pdf")},
            format="multipart",
        )
        self.assertEqual(subida.status_code, status.HTTP_200_OK)
        self.assertTrue(subida.data["cargada"])
        self.assertEqual(subida.data["nombre_archivo"], "contrato.pdf")

        mala = self.client.post(
            "/api/plantillas-contrato/internet-empresas/",
            {"archivo": SimpleUploadedFile("nota.txt", b"hola", content_type="text/plain")},
            format="multipart",
        )
        self.assertEqual(mala.status_code, status.HTTP_400_BAD_REQUEST)

        descarga = self.client.get("/api/plantillas-contrato/internet-empresas/archivo/")
        self.assertEqual(descarga.status_code, status.HTTP_200_OK)
        self.assertEqual(bytes(descarga.content), contenido)

    def test_generar_contrato_usa_plantilla_en_bd(self):
        import zipfile
        from io import BytesIO

        self._cargar_plantillas_en_bd()
        venta_id = self._venta_empresa_entel()
        response = self._post_generar_contrato(venta_id)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        with zipfile.ZipFile(BytesIO(response.content)) as archivo:
            nombres = set(archivo.namelist())
        self.assertIn("internet empresas.pdf", nombres)
        self.assertIn("bono duplica.pdf", nombres)


class AuthAndPermissionsTests(APITestCase):
    def test_login_and_me(self):
        User.objects.create_user("loginuser", password="secret123")
        response = self.client.post(
            "/api/auth/login/",
            {"username": "loginuser", "password": "secret123"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["username"], "loginuser")
        me = self.client.get("/api/auth/me/")
        self.assertEqual(me.status_code, status.HTTP_200_OK)

    def test_change_password(self):
        User.objects.create_user("claveuser", password="secret123")
        self.client.post(
            "/api/auth/login/",
            {"username": "claveuser", "password": "secret123"},
            format="json",
        )
        bad = self.client.post(
            "/api/auth/change-password/",
            {
                "current_password": "otra",
                "new_password": "Montaña-Verde-44",
                "confirm_password": "Montaña-Verde-44",
            },
            format="json",
        )
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

        mismatch = self.client.post(
            "/api/auth/change-password/",
            {
                "current_password": "secret123",
                "new_password": "Montaña-Verde-44",
                "confirm_password": "otraClave9",
            },
            format="json",
        )
        self.assertEqual(mismatch.status_code, status.HTTP_400_BAD_REQUEST)

        ok = self.client.post(
            "/api/auth/change-password/",
            {
                "current_password": "secret123",
                "new_password": "Montaña-Verde-44",
                "confirm_password": "Montaña-Verde-44",
            },
            format="json",
        )
        self.assertEqual(ok.status_code, status.HTTP_200_OK)
        me = self.client.get("/api/auth/me/")
        self.assertEqual(me.status_code, status.HTTP_200_OK)
        self.assertEqual(me.data["username"], "claveuser")

    def test_anonymous_cannot_list_ventas(self):
        response = self.client.get("/api/ventas/")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_asesor_cannot_patch_codigos_nor_estado(self):
        grupo = Group.objects.get(name="Asesor")
        ana = User.objects.create_user("ana_codigos", password="secret")
        ana.groups.add(grupo)

        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 50", velocidad=50, precio=40)
        flujo = Flujo.objects.create(nombre="Flujo codigos", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        self.client.force_authenticate(ana)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        venta_id = created.data["id"]

        forbidden_codigos = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"psi": "psi-11", "siro": "siro-22"},
            format="json",
        )
        self.assertEqual(forbidden_codigos.status_code, status.HTTP_403_FORBIDDEN)

        forbidden_estado = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"estado": "ANULADO"},
            format="json",
        )
        self.assertEqual(forbidden_estado.status_code, status.HTTP_403_FORBIDDEN)

        detalle = self.client.get(f"/api/ventas/{venta_id}/")
        self.assertEqual(detalle.status_code, status.HTTP_200_OK)
        self.assertEqual(detalle.data["psi"], "")

    def test_operaciones_can_patch_codigos_but_not_estado(self):
        admin = User.objects.create_superuser("admin_codigos", "admin_codigos@test.com", "pass")
        marta = User.objects.create_user("marta_codigos", password="secret")
        marta.groups.add(Group.objects.get(name="Operaciones"))

        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 50", velocidad=50, precio=40)
        flujo = Flujo.objects.create(nombre="Flujo codigos", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        self.client.force_authenticate(admin)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        venta_id = created.data["id"]
        Venta.objects.filter(id=venta_id).update(creado_por=marta)

        self.client.force_authenticate(marta)
        patched = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {
                "psi": "psi-11",
                "siro": "siro-22",
                "numero_oportunidad": "opp-33",
                "numero_orden": "ord-44",
                "oit": "oit-55",
                "cotizacion": "cot-66",
                "contrato": "ct-77",
                "numero_fijo": "01444555",
            },
            format="json",
        )
        self.assertEqual(patched.status_code, status.HTTP_200_OK)
        self.assertEqual(patched.data["psi"], "PSI-11")
        self.assertEqual(patched.data["siro"], "SIRO-22")
        self.assertEqual(patched.data["numero_oportunidad"], "OPP-33")
        self.assertEqual(patched.data["numero_orden"], "ORD-44")
        self.assertEqual(patched.data["oit"], "OIT-55")
        self.assertEqual(patched.data["cotizacion"], "COT-66")
        self.assertEqual(patched.data["contrato"], "CT-77")
        self.assertEqual(patched.data["numero_fijo"], "01444555")

        patched_saf = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"saf": "309341"},
            format="json",
        )
        self.assertEqual(patched_saf.status_code, status.HTTP_200_OK)
        self.assertEqual(patched_saf.data["saf"], "309341")

        forbidden = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"estado": "ANULADO"},
            format="json",
        )
        self.assertEqual(forbidden.status_code, status.HTTP_403_FORBIDDEN)

    def test_supervisor_sees_owner_and_can_reassign_venta(self):
        ana = User.objects.create_user("ana_reasignar", password="secret", first_name="Ana")
        ana.groups.add(Group.objects.get(name="Asesor"))
        luis = User.objects.create_user("luis_reasignar", password="secret", first_name="Luis")
        luis.groups.add(Group.objects.get(name="Asesor"))
        supervisor = User.objects.create_user("super_reasignar", password="secret")
        supervisor.groups.add(Group.objects.get(name="Supervisor"))
        marta = User.objects.create_user("marta_reasignar", password="secret")
        marta.groups.add(Group.objects.get(name="Operaciones"))

        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 50", velocidad=50, precio=40)
        flujo = Flujo.objects.create(nombre="Flujo reasignar", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        self.client.force_authenticate(ana)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        venta_id = created.data["id"]
        self.assertEqual(created.data["creado_por"], ana.id)
        self.assertEqual(created.data["creado_por_detalle"]["username"], "ana_reasignar")

        forbidden_list = self.client.get("/api/asesores/")
        self.assertEqual(forbidden_list.status_code, status.HTTP_403_FORBIDDEN)

        stolen = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"creado_por": luis.id},
            format="json",
        )
        self.assertEqual(stolen.status_code, status.HTTP_403_FORBIDDEN)

        self.client.force_authenticate(supervisor)
        listed = self.client.get("/api/ventas/")
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        sale = next(item for item in listed.data if item["id"] == venta_id)
        self.assertEqual(sale["creado_por_detalle"]["first_name"], "Ana")

        asesores = self.client.get("/api/asesores/")
        self.assertEqual(asesores.status_code, status.HTTP_200_OK)
        usernames = {item["username"] for item in asesores.data}
        self.assertIn("ana_reasignar", usernames)
        self.assertIn("luis_reasignar", usernames)
        self.assertNotIn("marta_reasignar", usernames)

        reassigned = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"creado_por": luis.id},
            format="json",
        )
        self.assertEqual(reassigned.status_code, status.HTTP_200_OK)
        self.assertEqual(reassigned.data["creado_por"], luis.id)
        self.assertEqual(reassigned.data["creado_por_detalle"]["username"], "luis_reasignar")

        rejected = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"creado_por": marta.id},
            format="json",
        )
        self.assertEqual(rejected.status_code, status.HTTP_400_BAD_REQUEST)

        self.client.force_authenticate(marta)
        marta_list = self.client.get("/api/ventas/")
        self.assertEqual(marta_list.status_code, status.HTTP_200_OK)
        self.assertTrue(any(item["id"] == venta_id for item in marta_list.data))

        marta_reasign = self.client.patch(
            f"/api/ventas/{venta_id}/",
            {"creado_por": ana.id},
            format="json",
        )
        self.assertEqual(marta_reasign.status_code, status.HTTP_200_OK)
        self.assertEqual(marta_reasign.data["creado_por"], ana.id)

    def test_asesor_and_operaciones_can_comment_on_historial(self):
        grupo = Group.objects.get(name="Asesor")
        ana = User.objects.create_user("ana_nota", password="secret", first_name="Ana")
        ana.groups.add(grupo)
        marta = User.objects.create_user("marta_nota", password="secret", first_name="Marta")
        marta.groups.add(Group.objects.get(name="Operaciones"))

        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 50", velocidad=50, precio=40)
        flujo = Flujo.objects.create(nombre="Flujo notas", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        self.client.force_authenticate(ana)
        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertEqual(created.data["comentarios"], [])
        venta_id = created.data["id"]

        empty = self.client.post(
            "/api/venta-comentarios/",
            {"venta": venta_id, "texto": "   "},
            format="json",
        )
        self.assertEqual(empty.status_code, status.HTTP_400_BAD_REQUEST)

        ana_nota = self.client.post(
            "/api/venta-comentarios/",
            {"venta": venta_id, "texto": "El cliente no contesta al agendamiento."},
            format="json",
        )
        self.assertEqual(ana_nota.status_code, status.HTTP_201_CREATED)
        self.assertEqual(ana_nota.data["texto"], "El cliente no contesta al agendamiento.")
        self.assertEqual(ana_nota.data["creado_por"], ana.id)

        Venta.objects.filter(id=venta_id).update(creado_por=marta)
        self.client.force_authenticate(marta)
        marta_nota = self.client.post(
            "/api/venta-comentarios/",
            {"venta": venta_id, "texto": "Observado por las firmas."},
            format="json",
        )
        self.assertEqual(marta_nota.status_code, status.HTTP_201_CREATED)

        detalle = self.client.get(f"/api/ventas/{venta_id}/")
        self.assertEqual(detalle.status_code, status.HTTP_200_OK)
        textos = [item["texto"] for item in detalle.data["comentarios"]]
        self.assertEqual(
            textos,
            [
                "El cliente no contesta al agendamiento.",
                "Observado por las firmas.",
            ],
        )

        luis = User.objects.create_user("luis_nota", password="secret")
        luis.groups.add(grupo)
        self.client.force_authenticate(luis)
        forbidden = self.client.post(
            "/api/venta-comentarios/",
            {"venta": venta_id, "texto": "No debería poder."},
            format="json",
        )
        self.assertEqual(forbidden.status_code, status.HTTP_400_BAD_REQUEST)

    def test_asesor_cannot_delete_flujo_paso(self):
        group = Group.objects.get(name="Asesor")
        user = User.objects.create_user("ana", password="secret")
        user.groups.add(group)
        self.client.force_authenticate(user)

        flujo = Flujo.objects.create(nombre="Flujo", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        flujo_paso = FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        response = self.client.delete(f"/api/flujo-pasos/{flujo_paso.id}/")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_create_role_and_user(self):
        admin = User.objects.create_superuser("roleadmin", "role@test.com", "pass")
        self.client.force_authenticate(admin)

        catalog = self.client.get("/api/auth/permissions/")
        self.assertEqual(catalog.status_code, status.HTTP_200_OK)
        self.assertTrue(catalog.data)

        created = self.client.post(
            "/api/auth/roles/",
            {"name": "Auditor", "permissions": ["api.view_venta", "api.view_all_ventas"]},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertIn("api.view_all_ventas", created.data["permissions"])

        user = self.client.post(
            "/api/auth/users/",
            {
                "username": "auditor1",
                "password": "secret123",
                "first_name": "Ada",
                "groups": [created.data["id"]],
            },
            format="json",
        )
        self.assertEqual(user.status_code, status.HTTP_201_CREATED)
        self.assertEqual(user.data["groups"], ["Auditor"])
        self.assertTrue(user.data["is_active"])

    def test_rename_and_delete_user(self):
        admin = User.objects.create_superuser("renameadmin", "r@test.com", "pass")
        other = User.objects.create_user("viejo", password="secret")
        self.client.force_authenticate(admin)

        renamed = self.client.patch(
            f"/api/auth/users/{other.id}/",
            {"username": "nuevo", "first_name": "Ana"},
            format="json",
        )
        self.assertEqual(renamed.status_code, status.HTTP_200_OK)
        self.assertEqual(renamed.data["username"], "nuevo")
        self.assertEqual(renamed.data["first_name"], "Ana")

        blocked = self.client.delete(f"/api/auth/users/{admin.id}/")
        self.assertEqual(blocked.status_code, status.HTTP_400_BAD_REQUEST)

        deleted = self.client.delete(f"/api/auth/users/{other.id}/")
        self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)

    def test_rename_flujo_and_delete_venta(self):
        admin = User.objects.create_superuser("flujoadmin", "f@test.com", "pass")
        self.client.force_authenticate(admin)
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        producto = Producto.objects.create(nombre="Fibra 10", velocidad=10, precio=20)
        flujo = Flujo.objects.create(nombre="Original", tipo_cliente=TipoCliente.PERSONA)
        paso = Paso.objects.create(nombre="Uno", descripcion="Uno")
        FlujoPaso.objects.create(flujo=flujo, paso=paso, orden=1)

        renamed = self.client.patch(
            f"/api/flujos/{flujo.id}/",
            {"nombre": "Flujo persona"},
            format="json",
        )
        self.assertEqual(renamed.status_code, status.HTTP_200_OK)
        self.assertEqual(renamed.data["nombre"], "Flujo persona")

        paso_renamed = self.client.patch(
            f"/api/pasos/{paso.id}/",
            {"nombre": "Validación"},
            format="json",
        )
        self.assertEqual(paso_renamed.status_code, status.HTTP_200_OK)
        self.assertEqual(paso_renamed.data["nombre"], "Validación")

        created = self.client.post(
            "/api/ventas/",
            {"cliente": cliente.id, "producto": producto.id, "flujo": flujo.id},
            format="json",
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)

        blocked_flujo = self.client.delete(f"/api/flujos/{flujo.id}/")
        self.assertEqual(blocked_flujo.status_code, status.HTTP_400_BAD_REQUEST)

        deleted = self.client.delete(f"/api/ventas/{created.data['id']}/")
        self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)

        deleted_flujo = self.client.delete(f"/api/flujos/{flujo.id}/")
        self.assertEqual(deleted_flujo.status_code, status.HTTP_204_NO_CONTENT)


SUNAT_EMPRESA_HTML = """
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Número de RUC:</h4>
    </div>
    <div class="col-sm-7">
      <h4 class="list-group-item-heading">20522317285 - INDOTECH SAC</h4>
    </div>
  </div>
</div>
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Domicilio Fiscal:</h4>
    </div>
    <div class="col-sm-7">
      <p class="list-group-item-text">CAL.LUIGGI BARSATO NRO. 167 LIMA - LIMA - SAN BORJA</p>
    </div>
  </div>
</div>
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Actividad(es) Económica(s):</h4>
    </div>
    <div class="col-sm-7">
      <table class="table tblResultado">
        <tbody>
          <tr><td>Principal - 6201 - ACTIVIDADES DE PROGRAMACION INFORMATICA</td></tr>
          <tr><td>Secundaria 1 - 6202 - CONSULTORES EN INFORMATICA</td></tr>
        </tbody>
      </table>
    </div>
  </div>
</div>
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Comprobantes de Pago c/aut. de impresión (F. 806 u 816):</h4>
    </div>
    <div class="col-sm-7">
      <table class="table tblResultado">
        <tbody>
          <tr><td>FACTURA</td></tr>
          <tr><td>BOLETA DE VENTA</td></tr>
        </tbody>
      </table>
    </div>
  </div>
</div>
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Sistema de Emisión Electrónica:</h4>
    </div>
    <div class="col-sm-7">
      <table class="table tblResultado">
        <tbody>
          <tr><td>FACTURA PORTAL DESDE 23/10/2015</td></tr>
          <tr><td>BOLETA PORTAL DESDE 16/11/2015</td></tr>
        </tbody>
      </table>
    </div>
  </div>
</div>
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Padrones:</h4>
    </div>
    <div class="col-sm-7">
      <table class="table tblResultado">
        <tbody>
          <tr><td>NINGUNO</td></tr>
        </tbody>
      </table>
    </div>
  </div>
</div>
<!-- <div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Razón Social:</h4>
    </div>
    <div class="col-sm-7">
      <p class="list-group-item-text">eeee</p>
    </div>
  </div>
</div> -->
"""

SUNAT_PERSONA_HTML = """
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Número de RUC:</h4>
    </div>
    <div class="col-sm-7">
      <h4 class="list-group-item-heading">10123456789 - PEREZ LOPEZ, JUAN CARLOS</h4>
    </div>
  </div>
</div>
"""

SUNAT_CE_HTML = """
<div class="list-group-item">
  <div class="row">
    <div class="col-sm-5">
      <h4 class="list-group-item-heading">Número de RUC:</h4>
    </div>
    <div class="col-sm-7">
      <h4 class="list-group-item-heading">15123456789 - SMITH JOHNSON MARIA ELENA</h4>
    </div>
  </div>
</div>
"""


class LookupRucTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("tester", "tester@test.com", "pass")
        self.client.force_authenticate(self.user)

    def test_parse_sunat_html_empresa_and_persona(self):
        from .sunat import parse_sunat_html

        empresa = parse_sunat_html(SUNAT_EMPRESA_HTML)
        self.assertEqual(empresa["ruc"], "20522317285")
        self.assertEqual(empresa["razon_social"], "INDOTECH SAC")
        self.assertEqual(empresa["tipo_cliente"], "EMPRESA")
        self.assertEqual(empresa["distrito"], "SAN BORJA")
        self.assertEqual(empresa["numero"], "167")
        self.assertEqual(empresa["tipo_direccion"], "CALLE")
        labels = [item["label"] for item in empresa["ficha"]]
        self.assertIn("Actividad(es) Económica(s)", labels)
        self.assertIn("Comprobantes de Pago c/aut. de impresión (F. 806 u 816)", labels)
        self.assertIn("Sistema de Emisión Electrónica", labels)
        self.assertIn("Padrones", labels)
        self.assertNotIn("Razón Social", labels)
        actividad = next(i for i in empresa["ficha"] if i["label"] == "Actividad(es) Económica(s)")
        self.assertIn("PRINCIPAL - 6201", actividad["value"])
        self.assertIn("SECUNDARIA 1 - 6202", actividad["value"])
        padrones = next(i for i in empresa["ficha"] if i["label"] == "Padrones")
        self.assertEqual(padrones["value"], "NINGUNO")

        persona = parse_sunat_html(SUNAT_PERSONA_HTML)
        self.assertEqual(persona["tipo_cliente"], "PERSONA")
        self.assertEqual(persona["tipo_documento"], "DNI")
        self.assertEqual(persona["numero_documento"], "12345678")
        self.assertEqual(persona["nombres"], "JUAN CARLOS")
        self.assertEqual(persona["apellidos"], "PEREZ LOPEZ")

        ce = parse_sunat_html(SUNAT_CE_HTML)
        self.assertEqual(ce["tipo_cliente"], "PERSONA")
        self.assertEqual(ce["tipo_documento"], "CE")
        self.assertEqual(ce["numero_documento"], "123456789")
        self.assertEqual(ce["nombres"], "MARIA ELENA")
        self.assertEqual(ce["apellidos"], "SMITH JOHNSON")

        encoded = parse_sunat_html(
            """
            <div class="list-group-item">
              <div class="row">
                <div class="col-sm-5">
                  <h4 class="list-group-item-heading">N&uacute;mero de RUC:</h4>
                </div>
                <div class="col-sm-7">
                  <h4 class="list-group-item-heading">20131312955 - SUPERINTENDENCIA NACIONAL DE ADUANAS Y DE ADMINISTRACION TRIBUTARIA - SUNAT</h4>
                </div>
              </div>
            </div>
            <div class="list-group-item">
              <div class="row">
                <div class="col-sm-5">
                  <h4 class="list-group-item-heading">Domicilio Fiscal:</h4>
                </div>
                <div class="col-sm-7">
                  <p class="list-group-item-text">AV. GARCILASO DE LA VEGA NRO. 1472 LIMA - LIMA - LIMA</p>
                </div>
              </div>
            </div>
            """
        )
        self.assertEqual(encoded["ruc"], "20131312955")
        self.assertIn("SUPERINTENDENCIA NACIONAL", encoded["razon_social"])
        self.assertEqual(encoded["tipo_direccion"], "AVENIDA")
        self.assertEqual(encoded["numero"], "1472")
        self.assertEqual(encoded["distrito"], "LIMA")

    def test_lookup_uses_existing_cliente_and_last_sale(self):
        representante = Persona.objects.create(
            cliente=Cliente.objects.create(tipo=TipoCliente.PERSONA),
            tipo_documento="DNI",
            numero_documento="87654321",
            nombres="Ana",
            apellidos="Perez",
            celular="987654321",
        )
        empresa_cliente = Cliente.objects.create(tipo=TipoCliente.EMPRESA)
        Empresa.objects.create(
            cliente=empresa_cliente,
            ruc="20522317285",
            razon_social="INDOTECH SAC",
            representante_legal=representante,
        )
        Direccion.objects.create(
            cliente=empresa_cliente,
            tipo="CALLE",
            direccion="Luiggi Barsato",
            numero="167",
            distrito="San Borja",
        )
        producto = Producto.objects.create(
            nombre="Fibra", velocidad=200, precio=79, tipo_cliente=TipoCliente.EMPRESA
        )
        flujo = Flujo.objects.create(nombre="Empresa", tipo_cliente=TipoCliente.EMPRESA)
        promo = Promocion.objects.create(nombre="Promo", descripcion="x")
        venta = self.client.post(
            "/api/ventas/",
            {
                "cliente": empresa_cliente.id,
                "producto": producto.id,
                "flujo": flujo.id,
                "promociones": [promo.id],
            },
            format="json",
        )
        self.assertEqual(venta.status_code, status.HTTP_201_CREATED)

        with patch("api.lookup_views.consultar_sunat") as mocked:
            response = self.client.get("/api/lookup/ruc/", {"ruc": "20522317285"})
            mocked.assert_not_called()

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["source"], "cliente")
        self.assertEqual(response.data["cliente_id"], empresa_cliente.id)
        self.assertEqual(response.data["razon_social"], "INDOTECH SAC")
        self.assertEqual(response.data["nombres"], "Ana")
        self.assertEqual(response.data["producto"], str(producto.id))
        self.assertEqual(response.data["promociones"], [str(promo.id)])

    def test_lookup_falls_back_to_sunat(self):
        with patch(
            "api.lookup_views.consultar_sunat",
            return_value={
                "ruc": "20522317285",
                "razon_social": "INDOTECH SAC",
                "tipo_cliente": "EMPRESA",
                "tipo_direccion": "CALLE",
                "direccion": "Luiggi Barsato",
                "numero": "167",
                "distrito": "SAN BORJA",
            },
        ):
            response = self.client.get("/api/lookup/ruc/", {"ruc": "20522317285"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["source"], "sunat")
        self.assertEqual(response.data["razon_social"], "INDOTECH SAC")
        self.assertIsNone(response.data.get("cliente_id"))

    def test_parse_representantes_html(self):
        from .sunat import parse_representantes

        html = """
        <table class="table">
          <thead>
            <tr>
              <th>Documento</th>
              <th>Nro. Documento</th>
              <th>Nombre</th>
              <th>Cargo</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td>DNI</td>
              <td>12345678</td>
              <td>PEREZ LOPEZ JUAN CARLOS</td>
              <td>GERENTE GENERAL</td>
              <td>05/06/2017</td>
            </tr>
            <tr>
              <td>DNI</td>
              <td>87654321</td>
              <td>RAMOS DIAZ ANA MARIA</td>
              <td>APODERADO</td>
              <td>01/02/2019</td>
            </tr>
          </tbody>
        </table>
        """
        reps = parse_representantes(html)
        self.assertEqual(len(reps), 2)
        self.assertEqual(reps[0]["tipo_documento"], "DNI")
        self.assertEqual(reps[0]["numero_documento"], "12345678")
        self.assertEqual(reps[0]["nombres"], "JUAN CARLOS")
        self.assertEqual(reps[0]["apellidos"], "PEREZ LOPEZ")
        self.assertEqual(reps[0]["cargo"], "GERENTE GENERAL")
        self.assertEqual(reps[0]["fecha_desde"], "05/06/2017")
        self.assertEqual(reps[1]["fecha_desde"], "01/02/2019")

    def test_lookup_sunat_single_representante_fills_form(self):
        with patch(
            "api.lookup_views.consultar_sunat",
            return_value={
                "ruc": "20123456789",
                "razon_social": "EMPRESA SAC",
                "tipo_cliente": "EMPRESA",
                "representantes": [
                    {
                        "tipo_documento": "DNI",
                        "numero_documento": "12345678",
                        "nombres": "JUAN CARLOS",
                        "apellidos": "PEREZ LOPEZ",
                        "nombre_completo": "PEREZ LOPEZ JUAN CARLOS",
                        "cargo": "GERENTE GENERAL",
                    }
                ],
            },
        ):
            response = self.client.get("/api/lookup/ruc/", {"ruc": "20123456789"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["nombres"], "JUAN CARLOS")
        self.assertNotIn("representantes", response.data)

    def test_lookup_sunat_multiple_representantes_are_listed(self):
        with patch(
            "api.lookup_views.consultar_sunat",
            return_value={
                "ruc": "20123456789",
                "razon_social": "EMPRESA SAC",
                "tipo_cliente": "EMPRESA",
                "representantes": [
                    {
                        "tipo_documento": "DNI",
                        "numero_documento": "12345678",
                        "nombres": "JUAN",
                        "apellidos": "PEREZ",
                        "nombre_completo": "PEREZ JUAN",
                        "cargo": "GERENTE",
                    },
                    {
                        "tipo_documento": "DNI",
                        "numero_documento": "87654321",
                        "nombres": "ANA",
                        "apellidos": "RAMOS",
                        "nombre_completo": "RAMOS ANA",
                        "cargo": "APODERADO",
                    },
                ],
            },
        ):
            response = self.client.get("/api/lookup/ruc/", {"ruc": "20123456789"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["representantes"]), 2)
        self.assertNotIn("nombres", response.data)

    def test_persona_and_direccion_are_saved_uppercase(self):
        persona = self.client.post(
            "/api/personas/",
            {
                "tipo_documento": "DNI",
                "numero_documento": "12345678",
                "nombres": "Ana María",
                "apellidos": "Pérez",
                "distrito_nacimiento": "lima",
                "padre": "carlos",
                "madre": "maria",
                "celular": "987654321",
                "correo": "ana.perez@example.com",
            },
            format="json",
        )
        self.assertEqual(persona.status_code, status.HTTP_201_CREATED)
        self.assertEqual(persona.data["nombres"], "ANA MARÍA")
        self.assertEqual(persona.data["apellidos"], "PÉREZ")

        direccion = self.client.post(
            "/api/direcciones/",
            {
                "cliente": persona.data["cliente"],
                "tipo": "CALLE",
                "direccion": "luiggi barsato",
                "numero": "167",
                "distrito": "san borja",
            },
            format="json",
        )
        self.assertEqual(direccion.status_code, status.HTTP_201_CREATED)
        self.assertEqual(direccion.data["direccion"], "LUIGGI BARSATO")
        self.assertEqual(direccion.data["distrito"], "SAN BORJA")

        updated = self.client.patch(
            f"/api/direcciones/{direccion.data['id']}/",
            {"distrito": "san isidro", "numero": "200"},
            format="json",
        )
        self.assertEqual(updated.status_code, status.HTTP_200_OK)
        self.assertEqual(updated.data["distrito"], "SAN ISIDRO")
        self.assertEqual(updated.data["numero"], "200")

        deleted = self.client.delete(f"/api/direcciones/{direccion.data['id']}/")
        self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)

    def test_duplicate_direccion_for_cliente_is_reused(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        payload = {
            "cliente": cliente.id,
            "tipo": "CALLE",
            "direccion": "luiggi barsato",
            "numero": "167",
            "distrito": "san borja",
        }
        first = self.client.post("/api/direcciones/", payload, format="json")
        second = self.client.post("/api/direcciones/", payload, format="json")
        self.assertEqual(first.status_code, status.HTTP_201_CREATED)
        self.assertEqual(second.status_code, status.HTTP_201_CREATED)
        self.assertEqual(first.data["id"], second.data["id"])
        self.assertEqual(Direccion.objects.filter(cliente=cliente).count(), 1)

        other = self.client.post(
            "/api/direcciones/",
            {**payload, "tienda": "A-12", "piso": "2", "galeria": "centro lima", "interior": "2"},
            format="json",
        )
        self.assertEqual(other.status_code, status.HTTP_201_CREATED)
        self.assertNotEqual(other.data["id"], first.data["id"])
        self.assertEqual(Direccion.objects.filter(cliente=cliente).count(), 2)

    def test_lookup_direccion_returns_saved_and_parsed(self):
        cliente = Cliente.objects.create(tipo=TipoCliente.PERSONA)
        Direccion.objects.create(
            cliente=cliente,
            tipo="CALLE",
            direccion="LUIGGI BARSATO",
            numero="167",
            distrito="SAN BORJA",
        )
        Direccion.objects.create(
            cliente=cliente,
            tipo="AVENIDA",
            direccion="JAVIER PRADO",
            numero="100",
            distrito="SAN ISIDRO",
        )
        saved = self.client.get("/api/lookup/direccion/", {"q": "luiggi"})
        self.assertEqual(saved.status_code, status.HTTP_200_OK)
        self.assertEqual(saved.data["items"][0]["direccion"], "LUIGGI BARSATO")

        parsed = self.client.get(
            "/api/lookup/direccion/",
            {"q": "CAL.LUIGGI BARSATO NRO. 167 LIMA - LIMA - SAN BORJA"},
        )
        self.assertEqual(parsed.data["parsed"]["tipo_direccion"], "CALLE")
        self.assertEqual(parsed.data["parsed"]["numero"], "167")
        self.assertEqual(parsed.data["parsed"]["distrito"], "SAN BORJA")


