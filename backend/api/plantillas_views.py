from io import BytesIO
from pathlib import Path

import pymupdf
from django.http import HttpResponse
from django.utils.http import content_disposition_header
from django.utils.text import slugify
from openpyxl import load_workbook
from rest_framework.decorators import api_view, permission_classes
from rest_framework.exceptions import NotFound, PermissionDenied, ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .contratos.mapas import campos_a_mapa, mapa_a_campos, json_a_mapa
from .models import Documento, PlantillaContrato, Producto, Promocion

MAX_PLANTILLA = 15 * 1024 * 1024
ACEPTA_OK = {".pdf", ".xlsx", ".xlsm", ".eml"}
TRANSFORMACIONES = (
    "doc_y_dni",
    "velocidades",
    "nombre_plan",
    "dia_mes_anio",
    "descuento",
    "descuento_con_igv",
    "precio_int_solo",
)


def _require_perm(user, perm, mensaje):
    if user.is_superuser or user.has_perm(perm):
        return
    raise PermissionDenied(mensaje)


def _can_write(user):
    return user.is_superuser or user.has_perm("api.change_plantillacontrato") or user.has_perm(
        "api.add_plantillacontrato"
    )


def _documento(clave: str) -> Documento:
    documento = (
        Documento.objects.filter(clave=clave)
        .prefetch_related("productos", "promociones")
        .first()
    )
    if documento is None:
        raise NotFound("Documento desconocido.")
    return documento


def _acepta(value) -> list[str]:
    if value in (None, ""):
        return [".pdf"]
    if isinstance(value, str):
        value = [value]
    if not isinstance(value, (list, tuple)):
        raise ValidationError({"acepta": "Usa una lista de extensiones, por ejemplo .pdf."})
    acepta = []
    for item in value:
        extension = str(item).strip().lower()
        if not extension:
            continue
        if not extension.startswith("."):
            extension = f".{extension}"
        if extension not in ACEPTA_OK:
            raise ValidationError({"acepta": f"Formato no válido: {extension}."})
        if extension not in acepta:
            acepta.append(extension)
    if not acepta:
        raise ValidationError({"acepta": "Indica al menos un formato."})
    return acepta


def _ids(value, campo: str) -> list[int] | None:
    if value is None:
        return None
    if not isinstance(value, (list, tuple)):
        raise ValidationError({campo: f"Usa una lista de {campo}."})
    ids = []
    for item in value:
        try:
            ids.append(int(item))
        except (TypeError, ValueError) as exc:
            raise ValidationError({campo: f"Hay un identificador inválido en {campo}."}) from exc
    return ids


def _clave_nueva(etiqueta: str) -> str:
    base = slugify(etiqueta)[:40].strip("-") or "documento"
    clave = base
    indice = 2
    while Documento.objects.filter(clave=clave).exists():
        suffix = f"-{indice}"
        clave = f"{base[: 40 - len(suffix)]}{suffix}"
        indice += 1
    return clave


def _mapa_desde_payload(data) -> dict | None:
    if "campos" in data:
        try:
            return campos_a_mapa(data.get("campos"))
        except (TypeError, ValueError) as exc:
            raise ValidationError({"campos": str(exc)}) from exc
    if "mapa" in data:
        mapa = data.get("mapa") or {}
        if not isinstance(mapa, dict):
            raise ValidationError({"mapa": "Usa un objeto de páginas y campos."})
        try:
            return campos_a_mapa(mapa_a_campos(mapa))
        except (TypeError, ValueError) as exc:
            raise ValidationError({"mapa": str(exc)}) from exc
    return None


def _serialize(documento: Documento, registro: PlantillaContrato | None = None) -> dict:
    productos = list(documento.productos.all())
    promociones = list(documento.promociones.all())
    data = {
        "clave": documento.clave,
        "archivo": documento.archivo,
        "etiqueta": documento.etiqueta or documento.clave,
        "descripcion": documento.descripcion,
        "grupo": documento.grupo or "Otros",
        "acepta": documento.acepta or [".pdf"],
        "hojas": documento.hojas or [],
        "requerida": documento.requerida,
        "mapa": documento.mapa or {},
        "campos": mapa_a_campos(documento.mapa or {}),
        "cargada": False,
        "tamano": 0,
        "nombre_archivo": "",
        "actualizado": None,
        "actualizado_por": None,
        "productos": [producto.id for producto in productos],
        "promociones": [promocion.id for promocion in promociones],
        "productos_detalle": [
            {"id": producto.id, "nombre": producto.nombre, "velocidad": producto.velocidad}
            for producto in productos
        ],
        "promociones_detalle": [
            {"id": promocion.id, "nombre": promocion.nombre} for promocion in promociones
        ],
    }
    if registro is None:
        return data
    usuario = registro.actualizado_por
    data.update(
        {
            "cargada": True,
            "tamano": registro.tamano,
            "nombre_archivo": registro.nombre_archivo,
            "actualizado": registro.actualizado.isoformat() if registro.actualizado else None,
            "actualizado_por": (
                {
                    "id": usuario.id,
                    "username": usuario.username,
                    "first_name": usuario.first_name,
                    "last_name": usuario.last_name,
                }
                if usuario
                else None
            ),
        }
    )
    return data


def _registro(clave: str) -> PlantillaContrato | None:
    return PlantillaContrato.objects.filter(clave=clave).select_related("actualizado_por").first()


def _paginas_necesarias(documento: Documento) -> int:
    mapa = json_a_mapa(documento.mapa) if documento.mapa else {}
    if not mapa:
        return 1
    return max(mapa) + 1


def _validar_archivo(documento: Documento, nombre: str, contenido: bytes) -> None:
    extension = Path(nombre).suffix.lower()
    acepta = documento.acepta or [".pdf"]
    if extension not in acepta:
        raise ValidationError({"archivo": f"Usa un archivo {', '.join(acepta)}."})
    if len(contenido) > MAX_PLANTILLA:
        raise ValidationError({"archivo": "El archivo no puede superar 15 MB."})
    if not contenido:
        raise ValidationError({"archivo": "El archivo está vacío."})

    if extension == ".pdf":
        if not contenido.startswith(b"%PDF"):
            raise ValidationError({"archivo": "El archivo no es un PDF válido."})
        pdf = pymupdf.open(stream=contenido, filetype="pdf")
        try:
            needed = _paginas_necesarias(documento)
            if len(pdf) < needed:
                raise ValidationError(
                    {
                        "archivo": (
                            f'Este PDF tiene {len(pdf)} página(s); se esperan al menos {needed} '
                            f'para “{documento.etiqueta}”.'
                        )
                    }
                )
        finally:
            pdf.close()
        return

    if extension in {".xlsm", ".xlsx"}:
        if contenido[:2] != b"PK":
            raise ValidationError({"archivo": "El archivo no es un Excel válido."})
        try:
            wb = load_workbook(
                BytesIO(contenido),
                keep_vba=extension == ".xlsm",
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise ValidationError({"archivo": "No se pudo leer el Excel."}) from exc
        try:
            for hoja in documento.hojas or []:
                if hoja not in wb.sheetnames:
                    raise ValidationError({"archivo": f'Falta la hoja "{hoja}" en el Excel.'})
        finally:
            wb.close()
        return

    if extension == ".eml":
        try:
            texto = contenido.decode("utf-8")
        except UnicodeDecodeError:
            texto = contenido.decode("latin-1")
        if "subject:" not in texto.lower():
            raise ValidationError({"archivo": "El correo debe incluir un campo subject."})


def _asignar(documento: Documento, data) -> None:
    producto_ids = _ids(data.get("productos"), "productos")
    promocion_ids = _ids(data.get("promociones"), "promociones")
    if producto_ids is not None:
        productos = list(Producto.objects.filter(id__in=producto_ids))
        if len(productos) != len(set(producto_ids)):
            raise ValidationError({"productos": "Hay un producto que no existe."})
        documento.productos.set(productos)
    if promocion_ids is not None:
        promociones = list(Promocion.objects.filter(id__in=promocion_ids))
        if len(promociones) != len(set(promocion_ids)):
            raise ValidationError({"promociones": "Hay una promoción que no existe."})
        documento.promociones.set(promociones)


def _crear_documento(request):
    if not _can_write(request.user):
        raise PermissionDenied("No tienes permiso para crear documentos.")
    etiqueta = str(request.data.get("etiqueta") or "").strip()
    if not etiqueta:
        raise ValidationError({"etiqueta": "El nombre es obligatorio."})
    documento = Documento.objects.create(
        clave=_clave_nueva(etiqueta),
        etiqueta=etiqueta[:120],
        descripcion=str(request.data.get("descripcion") or "").strip(),
        grupo=str(request.data.get("grupo") or "Contratos").strip()[:60],
        acepta=_acepta(request.data.get("acepta")),
        hojas=request.data.get("hojas") or [],
        mapa=_mapa_desde_payload(request.data) or {},
    )
    _asignar(documento, request.data)
    return Response(_serialize(_documento(documento.clave), None), status=201)


def _actualizar_documento(request, clave):
    _require_perm(
        request.user,
        "api.change_plantillacontrato",
        "No tienes permiso para editar este documento.",
    )
    documento = Documento.objects.filter(clave=clave).first()
    if documento is None:
        raise NotFound("Documento desconocido.")
    data = request.data
    if "etiqueta" in data:
        etiqueta = str(data.get("etiqueta") or "").strip()
        if not etiqueta:
            raise ValidationError({"etiqueta": "El nombre es obligatorio."})
        documento.etiqueta = etiqueta[:120]
    if "descripcion" in data:
        documento.descripcion = str(data.get("descripcion") or "").strip()
    if "grupo" in data:
        documento.grupo = str(data.get("grupo") or "").strip()[:60]
    if "acepta" in data:
        documento.acepta = _acepta(data.get("acepta"))
    if "hojas" in data:
        hojas = data.get("hojas") or []
        if not isinstance(hojas, (list, tuple)):
            raise ValidationError({"hojas": "Usa una lista de hojas."})
        documento.hojas = [str(item) for item in hojas]
    mapa = _mapa_desde_payload(data)
    if mapa is not None:
        for item in mapa_a_campos(mapa):
            transform = item.get("transform")
            if transform and transform not in TRANSFORMACIONES:
                raise ValidationError(
                    {"campos": f"Fórmula no válida: {transform}."}
                )
        documento.mapa = mapa
    documento.save()
    _asignar(documento, data)
    return Response(_serialize(_documento(clave), _registro(clave)))


def _borrar_documento(request, clave):
    _require_perm(
        request.user,
        "api.delete_plantillacontrato",
        "No tienes permiso para eliminar este documento.",
    )
    documento = Documento.objects.filter(clave=clave).first()
    if documento is None:
        raise NotFound("Documento desconocido.")
    PlantillaContrato.objects.filter(clave=clave).delete()
    documento.delete()
    return Response(status=204)


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def plantillas_contrato(request):
    if request.method == "POST":
        return _crear_documento(request)
    _require_perm(
        request.user,
        "api.view_plantillacontrato",
        "No tienes permiso para ver las plantillas de contrato.",
    )
    registros = {
        item.clave: item
        for item in PlantillaContrato.objects.select_related("actualizado_por")
    }
    documentos = Documento.objects.prefetch_related("productos", "promociones").order_by(
        "grupo", "etiqueta", "clave"
    )
    return Response([_serialize(item, registros.get(item.clave)) for item in documentos])


@api_view(["POST", "PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def subir_plantilla_contrato(request, clave):
    if request.method == "PATCH":
        return _actualizar_documento(request, clave)
    if request.method == "DELETE":
        return _borrar_documento(request, clave)
    if not _can_write(request.user):
        raise PermissionDenied("No tienes permiso para cargar plantillas de contrato.")
    documento = _documento(clave)
    archivo = request.FILES.get("archivo")
    if archivo is None:
        raise ValidationError({"archivo": "Selecciona un archivo."})
    contenido = archivo.read()
    _validar_archivo(documento, archivo.name, contenido)
    nombre = Path(archivo.name).name[:200]
    registro, _created = PlantillaContrato.objects.update_or_create(
        clave=clave,
        defaults={
            "nombre_archivo": nombre,
            "content_type": (archivo.content_type or "")[:120],
            "tamano": len(contenido),
            "contenido": contenido,
            "actualizado_por": request.user,
        },
    )
    if not documento.archivo:
        documento.archivo = nombre
        documento.save(update_fields=["archivo"])
    return Response(_serialize(_documento(clave), registro))


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def descargar_plantilla_contrato(request, clave):
    _require_perm(
        request.user,
        "api.view_plantillacontrato",
        "No tienes permiso para ver las plantillas de contrato.",
    )
    documento = _documento(clave)
    registro = PlantillaContrato.objects.filter(clave=clave).first()
    if registro is None:
        raise NotFound("Aún no hay un archivo cargado para esta plantilla.")
    response = HttpResponse(
        bytes(registro.contenido),
        content_type=registro.content_type or "application/octet-stream",
    )
    response["Content-Disposition"] = content_disposition_header(
        True, registro.nombre_archivo or documento.archivo or f"{clave}.bin"
    )
    return response
